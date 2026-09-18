"""Testes dos certificados PPTX da tecnica (v1.35.0): NR-06/12/18/35.

Parte 1 (unitaria, sem Office): tokenizacao do prepare, substituicao de
tokens preservando runs, valores de data por extenso e deteccao de modelo.
Parte 2 (E2E, exige Microsoft PowerPoint): gera os 4 PDFs de verdade,
conferindo nome/CPF/data por extenso/numero do certificado dentro do PDF,
e salva copias em comparacao_pptx/ para conferencia manual.

Roda direto:  python test_pptx_certificados.py
"""

import re
import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from pptx import Presentation
from pptx.util import Pt

import src.core.pptx_certificate_service as pptx_cert
from tools.prepare_nr_pptx import ORIGEM, tokenizar, texto_completo

FALHAS = []
TMP = Path(tempfile.mkdtemp(prefix="test_pptx_cert_"))


def check(nome, ok):
    print(f"  [{'OK' if ok else 'FALHOU'}] {nome}")
    if not ok:
        FALHAS.append(nome)


def _pptx_com_runs(caminho: Path, runs):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 4000000, 1000000)
    p = box.text_frame.paragraphs[0]
    for t in runs:
        r = p.add_run()
        r.text = t
        r.font.size = Pt(20)
    prs.save(str(caminho))


def parte1_unitaria():
    print("\n--- Parte 1: unitaria (sem Office) ---")

    # valores de data por extenso
    emp = SimpleNamespace(nome="jose da silva", cpf="111.222.333-44")
    v = pptx_cert.valores_certificado(emp, date(2026, 9, 5))
    check("T1 dia sem zero a esquerda", v["DIA"] == "5")
    check("T2 mes por extenso minusculo", v["MES"] == "setembro")
    check("T3 nome em caixa alta", v["NOME"] == "JOSE DA SILVA")
    check("T4 CPF como cadastrado", v["CPF"] == "111.222.333-44")

    # substituicao preservando runs
    p_file = TMP / "runs.pptx"
    _pptx_com_runs(p_file, ["Certificamos que ", "{{NOME}}", ", CPF: ", "{{CPF}}",
                            " no dia ", "{{DIA}}", " de ", "{{MES}}", " de ", "{{ANO}}"])
    prs = Presentation(str(p_file))
    para = prs.slides[0].shapes[0].text_frame.paragraphs[0]
    ok = pptx_cert._substituir_no_paragrafo(para, {
        "NOME": "JOAO PEREIRA", "CPF": "999.888.777-66",
        "DIA": "1", "MES": "janeiro", "ANO": "2026",
    })
    txt = "".join(r.text for r in para.runs)
    check("T5 tokens substituidos em runs quebrados",
          ok and txt == "Certificamos que JOAO PEREIRA, CPF: 999.888.777-66"
                        " no dia 1 de janeiro de 2026")
    # a formatacao (tamanho 20pt) sobreviveu em todos os runs nao vazios
    tamanhos = {r.font.size.pt for r in para.runs if r.text}
    check("T6 formatacao do run preservada", tamanhos == {20.0})

    # token inexistente fica intacto
    p_file2 = TMP / "runs2.pptx"
    _pptx_com_runs(p_file2, ["Olá {{DESCONHECIDO}} e {{NOME}}"])
    prs2 = Presentation(str(p_file2))
    para2 = prs2.slides[0].shapes[0].text_frame.paragraphs[0]
    pptx_cert._substituir_no_paragrafo(para2, {"NOME": "MARIA"})
    txt2 = "".join(r.text for r in para2.runs)
    check("T7 token desconhecido intacto", txt2 == "Olá {{DESCONHECIDO}} e MARIA")

    # deteccao de modelo com get_templates_dir patchado
    original = pptx_cert.get_templates_dir
    try:
        pptx_cert.get_templates_dir = lambda: TMP / "vazio"
        check("T8 modelo ausente -> None",
              pptx_cert.get_pptx_template_path("NR-06") is None)
    finally:
        pptx_cert.get_templates_dir = original

    # tokenizacao do prepare a partir dos originais (idempotente)
    if ORIGEM.exists():
        alvo = TMP / "certificados_pptx"
        n_ok = 0
        for src in sorted(ORIGEM.glob("*.pptx")):
            m = re.search(r"NR\s*(\d+)", src.stem, re.IGNORECASE)
            if not m:
                continue
            nr = f"NR-{int(m.group(1)):02d}"
            tokenizar(src, alvo / f"{nr}.pptx")
            texto = texto_completo(alvo / f"{nr}.pptx")
            tem = all(t in texto for t in ("{{NOME}}", "{{CPF}}", "{{DIA}}", "{{MES}}", "{{ANO}}"))
            resto = re.findall(r"\b\d{1,2}\s+de\s+[a-zà-ú]+\s+de\s+\d{4}\b", texto, re.IGNORECASE)
            check(f"T9 {nr} tokenizado sem sobras", tem and not resto)
            n_ok += 1
        check("T10 quatro modelos tokenizados", n_ok == 4)
    else:
        print("  [SKIP] MODELOS NR pptx/ nao encontrado (T9/T10)")


def _office_ok() -> bool:
    try:
        import comtypes.client

        app = comtypes.client.CreateObject("PowerPoint.Application", dynamic=True)
        app.Quit()
        return True
    except Exception:
        return False


def parte2_e2e():
    print("\n--- Parte 2: E2E com PowerPoint (PDFs reais) ---")
    if not ORIGEM.exists():
        print("  [SKIP] MODELOS NR pptx/ nao encontrado")
        return
    originais = {re.search(r"NR\s*(\d+)", s.stem, re.IGNORECASE).group(1): s
                 for s in ORIGEM.glob("*.pptx") if re.search(r"NR\s*(\d+)", s.stem, re.IGNORECASE)}
    # aponta o servico para os modelos tokenizados em TMP
    pptx_cert.get_templates_dir = lambda: TMP
    if not _office_ok():
        print("  [SKIP] Microsoft PowerPoint nao disponivel (unitarios ja validados)")
        return

    comp_dir = Path(__file__).resolve().parent / "comparacao_pptx"
    comp_dir.mkdir(exist_ok=True)
    emp = SimpleNamespace(nome="Teste Da Silva", cpf="111.222.333-44")
    data = date(2026, 9, 15)

    import fitz

    for nr in ("NR-06", "NR-12", "NR-18", "NR-35"):
        if nr.split("-")[1] not in originais:
            continue
        saida = TMP / "saida" / f"{nr}_TESTE.pdf"
        try:
            pptx_cert.gerar_pdf_pptx(nr, emp, data, "CERT-099999", saida)
        except Exception as e:
            check(f"E {nr} gerou PDF ({e})", False)
            continue
        ok_pdf = saida.exists() and saida.read_bytes()[:5] == b"%PDF-"
        check(f"E {nr} PDF valido", ok_pdf)
        doc = fitz.open(str(saida))
        texto = " ".join(p.get_text() for p in doc)
        paginas = len(doc)
        doc.close()
        check(f"E {nr} nome no PDF", "TESTE DA SILVA" in texto.upper())
        check(f"E {nr} CPF no PDF", "111.222.333-44" in texto)
        check(f"E {nr} data por extenso", "15 de setembro de 2026" in texto)
        check(f"E {nr} numero do certificado", "CERT-099999" in texto)
        check(f"E {nr} 2 paginas (certificado + conteudo)", paginas == 2)
        shutil.copyfile(saida, comp_dir / f"{nr}_EXEMPLO.pdf")
        print(f"        copia: comparacao_pptx/{nr}_EXEMPLO.pdf")


def main():
    try:
        parte1_unitaria()
        parte2_e2e()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print()
    if FALHAS:
        print(f"FALHAS ({len(FALHAS)}): {', '.join(FALHAS)}")
        sys.exit(1)
    print("TESTES PPTX CERTIFICADOS OK")


if __name__ == "__main__":
    main()
