"""Testes de robustez das importações — v1.45.1.

Padrão standalone: `python test_import_robusto.py`.

Cobre:
- CPF na planilha: formatado (ponto/hífen), números puros, célula numérica
  do Excel (int e float x.0) — todos importam
- CPF com 12 dígitos / menos de 11 dígitos — rejeitado com mensagem clara
- Erro de banco em uma linha NÃO aborta o resto da planilha
- Telefone como célula numérica (float) é aceito
- Duplicado por nome é ignorado
- guards de float em aso_importer / blocking_importer / veiculo_importer
- portal: máscara de CPF no form e botão Exibir senha nos templates
- modelos xlsx com coluna CPF como Texto e exemplos em números puros
"""

import sys
import tempfile
from pathlib import Path

FALHAS = []


def check(nome, ok):
    print(f"[{'OK' if ok else 'FALHOU'}] {nome}")
    if not ok:
        FALHAS.append(nome)


def _xlsx_tmp(linhas):
    """Cria planilha tmp (header + linhas). Retorna caminho."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome", "CPF", "Funcao", "Telefone", "Data Nascimento",
               "Tipo Sanguineo", "Data Admissao", "Registro CTPS", "CNH EAR"])
    for l in linhas:
        ws.append(l)
    import tempfile
    fd = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    wb.save(fd.name)
    wb.close()
    return fd.name


def make_env():
    tmp = Path(tempfile.mkdtemp(prefix="imp_robusto_"))
    import src.core.employee_repo as er_mod
    import src.utils.funcoes_store as fs_mod
    er_mod.get_db_path = lambda: tmp / "certificados.db"
    fs_mod.get_data_dir = lambda: tmp
    return tmp


def limpar_repo():
    from src.core.employee_repo import EmployeeRepository
    repo = EmployeeRepository()
    for e in repo.get_all(limit=100000):
        repo.delete(e.id)
    return repo


def importar(linhas, repo=None):
    from src.utils.excel_importer import import_employees_from_excel
    from src.core.employee_repo import EmployeeRepository
    arq = _xlsx_tmp(linhas)
    try:
        return import_employees_from_excel(arq, repo or EmployeeRepository())
    finally:
        Path(arq).unlink(missing_ok=True)


def test_cpfs(repo):
    from src.core.employee_repo import EmployeeRepository

    imp, dup, err, det = importar([
        ["Joao Formatado", "529.982.247-25", "", "", "", "", "", "", ""],
        ["Maria Pura", "11144477735", "", "", "", "", "", "", ""],
        ["Carlos NumInt", 52998224725, "", "", "", "", "", "", ""],
        ["Daniela NumFloat", 11144477735.0, "", "", "", "", "", "", ""],
    ])
    check("CPF formatado/importado: 4 importados", imp == 4)
    check("sem erros", err == 0)
    e1 = repo.search("Joao Formatado", limit=1)[0]
    check("CPF formatado fica com máscara", e1.cpf == "529.982.247-25")
    e4 = repo.search("Daniela NumFloat", limit=1)[0]
    check("CPF float do Excel (x.0) importa sem perder dígito",
          e4.cpf == "111.444.777-35")


def test_rejeitados():
    imp, dup, err, det = importar([
        ["Doze Digitos", "529982247250", "", "", "", "", "", "", ""],
        ["Nove Digitos", "007007007", "", "", "", "", "", "", ""],
        ["Com Letras", "abc529", "", "", "", "", "", "", ""],
    ])
    check("CPF inválidos: 3 linhas rejeitadas", err == 3)
    check("CPF > 11 dígitos: mensagem clara",
          any("mais de 11" in d for d in det))
    check("CPF < 11 dígitos: dica de zeros à esquerda",
          any("menos de 11" in d for d in det))
    check("CPF com letras: também rejeitado (não importa ninguém)",
          imp == 0 and any("Com Letras" not in d for d in det) or imp == 0)


def test_erro_banco_nao_aborta(repo):
    from src.core.employee_repo import EmployeeRepository

    class RepoBomba:
        def __init__(self):
            self._real = EmployeeRepository()
            self.chamadas = 0

        def search(self, nome, limit=5):
            return self._real.search(nome, limit=limit)

        def create(self, *a, **k):
            self.chamadas += 1
            if self.chamadas == 2:
                raise RuntimeError("database is locked (simulado)")
            return self._real.create(*a, **k)

    imp, dup, err, det = importar([
        ["Ana Primeira", "52998224725", "", "", "", "", "", "", ""],
        ["Bruno Bomba", "11144477735", "", "", "", "", "", "", ""],
        ["Carla Terceira", "52998224725", "", "", "", "", "", "", ""],
    ], repo=RepoBomba())
    check("erro em 1 linha não aborta: 2 importadas mesmo assim", imp == 2)
    check("linha com falha virou erro com mensagem", err == 1
          and any("erro inesperado" in d and "Bruno Bomba" in d for d in det))


def test_telefone_float_e_duplicado():
    imp, dup, err, det = importar([
        ["Telefonista", "52998224725", "", 11999998888.0, "", "", "", "", ""],
        ["Telefonista", "52998224725", "", "", "", "", "", "", ""],
    ])
    check("telefone numérico (float) aceito", imp == 1 and err == 0)
    check("nome repetido vira duplicado", dup == 1)


def test_outros_importadores():
    from src.utils.aso_importer import _so_digitos
    from src.utils.blocking_importer import _digits
    from src.utils.veiculo_importer import _txt
    check("aso _so_digitos float ok", _so_digitos(52998224725.0) == "52998224725")
    check("aso _so_digitos formatado ok",
          _so_digitos("529.982.247-25") == "52998224725")
    check("blocking _digits float ok", _digits(11144477735.0) == "11144477735")
    check("veiculo _txt float ok", _txt(2020.0) == "2020" and _txt("Carro") == "Carro")


def test_templates_web():
    base = Path("src/web/templates")
    form = (base / "funcionario_form.html").read_text(encoding="utf-8")
    check("form funcionário tem máscara de CPF",
          'id="cpf"' in form and "replace(/\\D/g" in form and "slice(0, 11)" in form)
    check("form funcionário limita 11 dígitos", 'maxlength="14"' in form)
    usu = (base / "usuarios.html").read_text(encoding="utf-8")
    check("usuarios.html tem botão Exibir senha",
          "_alternarSenha" in usu and ">Exibir</button>" in usu)
    troca = (base / "troca_senha.html").read_text(encoding="utf-8")
    check("troca_senha.html tem botão Exibir senha",
          "_alternarSenha" in troca and troca.count(">Exibir</button>") == 3)


def test_modelos_xlsx():
    import openpyxl
    pasta = Path("MODELOS DE IMPORTACAO")
    wb = openpyxl.load_workbook(pasta / "MODELO FUNCIONARIOS.xlsx")
    ws = wb["FUNCIONARIOS"]
    check("MODELO FUNCIONARIOS: exemplo CPF em números puros",
          ws["B2"].value == "52998224725")
    check("MODELO FUNCIONARIOS: coluna CPF como Texto",
          ws["B2"].number_format == "@")
    check("MODELO FUNCIONARIOS: comentário no cabeçalho do CPF",
          bool(ws["B1"].comment) and "11 numeros" in ws["B1"].comment.text)
    check("MODELO FUNCIONARIOS: aba LEIA-ME presente", "LEIA-ME" in wb.sheetnames)
    wb.close()
    for nome in ("MODELO ASO.xlsx", "MODELO CARTOES BLOQUEIO.xlsx"):
        wb = openpyxl.load_workbook(pasta / nome)
        ws = wb[wb.sheetnames[0]]
        check(f"{nome}: CPF Texto + exemplo puro + comentário",
              ws["B2"].number_format == "@" and ws["B2"].value == "52998224725"
              and bool(ws["B1"].comment))
        wb.close()


def main():
    make_env()
    repo = limpar_repo()
    test_cpfs(repo)
    test_rejeitados()
    test_erro_banco_nao_aborta(repo)
    test_telefone_float_e_duplicado()
    test_outros_importadores()
    test_templates_web()
    test_modelos_xlsx()
    print()
    if FALHAS:
        print(f"FALHAS: {len(FALHAS)}")
        for f in FALHAS:
            print(f"  - {f}")
        sys.exit(1)
    print("Tudo OK")


if __name__ == "__main__":
    main()
