"""Portal Web — Gestão de Frota (ROADMAP 2.29, v1.33.0).

Veículos, empresas de veículos, fornecedores, documentos (pasta virtual),
laudos com vencimento, movimentações (saída/entrada) e solicitações de
abastecimento com PDF. Escrita exige admin/emissor; consulta navega em
modo leitura.
"""

import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote_plus

from markupsafe import Markup

from fastapi import Request, Form, UploadFile, File
from fastapi.responses import (
    RedirectResponse, Response, FileResponse,
)

from src.web import auth
from src.web.permissions import pode_escrever
from src.core.frota_repo import (
    FrotaRepository, TIPOS_VEICULO, SUBTIPOS_CAMINHAO, TIPOS_COMBUSTIVEL,
    TIPOS_LAUDO, CARROCERIAS, CHECKLIST_GRUPOS, CHECKLIST_DIAS,
    label_tipo, label_subtipo, label_combustivel,
    label_laudo, veiculo_rotulo, rotulo_revisao,
)

_PER_PAGE = 20
_PER_OPCOES = (10, 20, 25, 50)
_MIME = {"pdf": "application/pdf", "jpg": "image/jpeg", "jpeg": "image/jpeg",
         "png": "image/png", "gif": "image/gif", "txt": "text/plain"}

# 2.29.7: Arla/Diesel/Arla+Diesel não são combustíveis de carro de passeio
_TIPOS_SEM_DIESEL = {"carro"}
_COMB_BLOQUEADOS_LEVES = {"arla", "diesel", "arla_diesel"}

# 2.35.2: motivos de bloqueio de solicitação de abastecimento
_MOTIVOS_BLOQUEIO = ("Não usada", "Erro de lançamento", "Cancelado", "Outro")

# tags para documentos do veículo (2.29.7)
TAGS_DOC = [
    ("manutencao", "Manutenção (Nota Fiscal)"),
    ("documento", "Documento"),
    ("abastecimento", "Abastecimento"),
    ("outros", "Outros"),
]


def _br(iso) -> str:
    try:
        return datetime.strptime(str(iso)[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except Exception:
        return ""


def _iso(data_br: str) -> str:
    """dd/mm/aaaa -> aaaa-mm-dd (None se inválida)."""
    try:
        return datetime.strptime((data_br or "").strip(),
                                 "%d/%m/%Y").strftime("%Y-%m-%d")
    except ValueError:
        return None


def _cnpj_invalido(cnpj: str):
    """CNPJ opcional: se preenchido, exige 14 dígitos (formato, sem DV).

    Retorna a mensagem de erro ou None se ok (v1.45.3).
    """
    digitos = re.sub(r"\D", "", cnpj or "")
    if not digitos:
        return None
    if len(digitos) != 14 or digitos == digitos[0] * 14:
        return "CNPJ inválido: informe os 14 dígitos (ou deixe vazio)."
    return None


def _mime(tipo: str) -> str:
    return _MIME.get((tipo or "").lower(), "application/octet-stream")


def register(app, deps):
    templates = deps["templates"]
    ctx = deps["ctx"]
    flash = deps["flash"]
    users = deps["users"]

    def _audit(request: Request, acao: str, username: str,
               alvo: str = "", detalhe: str = ""):
        try:
            users.audit(acao, username, alvo, detalhe)
        except Exception:
            pass

    def _bloqueio(request: Request, user: dict, url: str):
        if pode_escrever(user["papel"], "frota"):
            return None
        flash(request, erro="Seu papel é somente leitura neste módulo.")
        return RedirectResponse(url, status_code=303)

    def _repo() -> FrotaRepository:
        return FrotaRepository()

    def _paginacao(request: Request, total: int, page: int, per: int = _PER_PAGE):
        total_paginas = max(1, (total + per - 1) // per)
        page = min(max(1, page), total_paginas)
        return page, total_paginas

    def _ler_extras(fdata) -> list:
        """Lê as linhas de itens extras do formulário (2.35.2):
        extra_desc_N / extra_qtd_N / extra_val_N."""
        extras = []
        for i in range(20):
            desc = str(fdata.get(f"extra_desc_{i}") or "").strip()
            qtd = str(fdata.get(f"extra_qtd_{i}") or "").strip()
            val = str(fdata.get(f"extra_val_{i}") or "").strip()
            if desc or qtd or val:
                extras.append({"desc": desc, "qtd": qtd, "valor": val})
        return extras

    def _svg_grafico_custo(serie: list, largura=720, altura=260):
        """Gráfico SVG offline (2.35.1 / 2.36): barras empilhadas (combustível
        azul + itens extras amarelo) e linha de litros (verde) dos últimos 12
        meses. Cada mês é um <g class="g-mes"> com os valores em data-*
        (tooltip interativo) e <title> nativo."""
        if not serie:
            return ""

        def _fmt(v):
            return f"{float(v):.2f}".replace(".", ",")

        maxv = max(max((s["combustivel"] + s["extras"]) for s in serie), 0.01)
        maxl = max(max(s["litros"] for s in serie), 0.01)
        m_e, m_d, m_t, m_b = 46, 10, 14, 28
        pw, ph = largura - m_e - m_d, altura - m_t - m_b
        n = len(serie)
        passo = pw / n
        bw = min(34.0, passo * 0.55)
        y_base = m_t + ph
        partes = [f'<line x1="{m_e}" y1="{y_base}" x2="{m_e + pw}" '
                  f'y2="{y_base}" stroke="#D6DEE8"/>']
        for i, s in enumerate(serie):
            xc = m_e + i * passo + passo / 2
            h_c = s["combustivel"] / maxv * ph
            h_x = s["extras"] / maxv * ph
            x = m_e + i * passo + (passo - bw) / 2
            g = [f'<g class="g-mes" data-mes="{s["rotulo"]}" '
                 f'data-comb="{_fmt(s["combustivel"])}" '
                 f'data-ext="{_fmt(s["extras"])}" '
                 f'data-litros="{_fmt(s["litros"])}">']
            g.append(f'<title>{s["rotulo"]}: Combustível R$ {_fmt(s["combustivel"])}'
                     f' · Extras R$ {_fmt(s["extras"])}'
                     f' · Litros {_fmt(s["litros"])}</title>')
            if h_c > 0:
                g.append(f'<rect class="b-comb" x="{x:.1f}" y="{y_base - h_c:.1f}" '
                         f'width="{bw:.1f}" height="{h_c:.1f}" '
                         f'fill="#2E6DA4"/>')
            if h_x > 0:
                g.append(f'<rect class="b-ext" x="{x:.1f}" '
                         f'y="{y_base - h_c - h_x:.1f}" '
                         f'width="{bw:.1f}" height="{h_x:.1f}" '
                         f'fill="#E6A23C"/>')
            g.append(f'<rect class="captura" x="{m_e + i * passo:.1f}" y="{m_t}" '
                     f'width="{passo:.1f}" height="{ph:.1f}" '
                     f'fill="transparent"/>')
            g.append(f'<text x="{xc:.1f}" y="{altura - 10}" font-size="9" '
                     f'fill="#5A6B7C" text-anchor="middle">'
                     f'{s["rotulo"]}</text>')
            g.append('</g>')
            partes.append("".join(g))
        pts = []
        for i, s in enumerate(serie):
            x = m_e + i * passo + passo / 2
            y = y_base - (s["litros"] / maxl) * ph
            pts.append((x, y))
        if any(s["litros"] for s in serie):
            partes.append('<polyline points="'
                          + " ".join(f"{x:.1f},{y:.1f}" for x, y in pts)
                          + '" fill="none" stroke="#256B28" stroke-width="2"/>')
            for x, y in pts:
                partes.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="2.5" '
                              f'fill="#256B28"/>')
        partes.append(f'<text x="{m_e - 6}" y="{m_t + 4}" font-size="9" '
                      f'fill="#5A6B7C" text-anchor="end">'
                      f'R$ {maxv:,.0f}</text>'.replace(",", "."))
        partes.append(f'<text x="{m_e - 6}" y="{y_base}" font-size="9" '
                      f'fill="#5A6B7C" text-anchor="end">0</text>')
        return Markup(
            f'<svg viewBox="0 0 {largura} {altura}" width="{largura}" '
            f'height="{altura}" role="img" '
            f'style="width:100%;height:auto;" '
            f'xmlns="http://www.w3.org/2000/svg">{"".join(partes)}</svg>')

    # paleta p/ gráficos multi-veículo (2.37.2)
    _PALETA = ("#2E6DA4", "#E6A23C", "#256B28", "#B03A5B", "#6A5ACD",
               "#00838F", "#8D6E63", "#546E7A", "#C0CA33", "#AD1457",
               "#00897B", "#7B1FA2")
    _COMB_CORES = {"gasolina": "#2E6DA4", "diesel": "#37474F",
                   "alcool": "#256B28", "gnv": "#E6A23C", "arla": "#7E57C2",
                   "arla_diesel": "#6D4C41", "outros": "#78909C"}

    def _svg_custos_todos(serie: list, largura=860, altura=300):
        """Gráfico mensal empilhado POR VEÍCULO (2.37.2): cada mês tem um
        segmento por veículo (comb+extras somados); tooltip lista o detalhe."""
        veic_ids, veic_rot = [], {}
        for s in serie:
            for vid, d in (s.get("por_veiculo") or {}).items():
                if vid not in veic_rot:
                    veic_ids.append(vid)
                    veic_rot[vid] = d["rotulo"]
        if not veic_ids:
            return ""

        def _fmt(v):
            return f"{float(v):.2f}".replace(".", ",")

        def _esc(t):
            return (str(t).replace("&", "&amp;").replace("<", "&lt;")
                    .replace(">", "&gt;").replace('"', "&quot;"))

        maxv = 0.01
        for s in serie:
            tot = sum(d["combustivel"] + d["extras"]
                      for d in (s.get("por_veiculo") or {}).values())
            maxv = max(maxv, tot)
        m_e, m_d, m_t, m_b = 52, 10, 14, 28
        pw, ph = largura - m_e - m_d, altura - m_t - m_b
        n = len(serie)
        passo = pw / n
        bw = min(38.0, passo * 0.6)
        y_base = m_t + ph
        partes = [f'<line x1="{m_e}" y1="{y_base}" x2="{m_e + pw}" '
                  f'y2="{y_base}" stroke="#D6DEE8"/>']
        for i, s in enumerate(serie):
            detalhes = []
            acum = 0.0
            x = m_e + i * passo + (passo - bw) / 2
            g = [f'<g class="g-mes" data-mes="{_esc(s["rotulo"])}">']
            for vid in veic_ids:
                d = (s.get("por_veiculo") or {}).get(vid)
                if not d:
                    continue
                tot_v = d["combustivel"] + d["extras"]
                if tot_v <= 0:
                    continue
                detalhes.append(
                    f'{_esc(d["rotulo"])}: R$ {_fmt(tot_v)} '
                    f'(comb R$ {_fmt(d["combustivel"])} · extras R$ {_fmt(d["extras"])})')
                h = tot_v / maxv * ph
                g.append(f'<rect x="{x:.1f}" y="{y_base - acum - h:.1f}" '
                         f'width="{bw:.1f}" height="{h:.1f}" '
                         f'fill="{_PALETA[veic_ids.index(vid) % len(_PALETA)]}"/>')
                acum += h
            g.append(f'<title>{_esc(s["rotulo"])} — ' +
                     ("; ".join(detalhes) if detalhes else "Sem custos") +
                     '</title>')
            g.append(f'<rect class="captura" data-mes="{_esc(s["rotulo"])}" '
                     f'data-det="{"||".join(detalhes) if detalhes else "Sem custos."}" '
                     f'x="{m_e + i * passo:.1f}" y="{m_t}" '
                     f'width="{passo:.1f}" height="{ph:.1f}" fill="transparent"/>')
            g.append(f'<text x="{m_e + i * passo + passo / 2:.1f}" '
                     f'y="{altura - 10}" font-size="9" fill="#5A6B7C" '
                     f'text-anchor="middle">{_esc(s["rotulo"])}</text>')
            g.append('</g>')
            partes.append("".join(g))
        partes.append(f'<text x="{m_e - 6}" y="{m_t + 4}" font-size="9" '
                      f'fill="#5A6B7C" text-anchor="end">'
                      f'R$ {maxv:,.0f}</text>'.replace(",", "."))
        partes.append(f'<text x="{m_e - 6}" y="{y_base}" font-size="9" '
                      f'fill="#5A6B7C" text-anchor="end">0</text>')
        # legenda (fora do svg, o template monta)
        leg = "".join(
            f'<span class="leg-item"><span class="leg-cor" '
            f'style="background:{_PALETA[i % len(_PALETA)]}"></span>'
            f'{_esc(veic_rot[vid])}</span>'
            for i, vid in enumerate(veic_ids))
        return Markup(
            f'<svg viewBox="0 0 {largura} {altura}" width="{largura}" '
            f'height="{altura}" role="img" '
            f'style="width:100%;height:auto;" '
            f'xmlns="http://www.w3.org/2000/svg">{"".join(partes)}</svg>'), leg

    def _svg_custos_comb(itens: list, largura=860, altura=300):
        """Gráfico POR VEÍCULO segmentado por tipo de combustível + extras
        (2.37.2/2.38.2): BARRAS HORIZONTAIS — nome do veículo à esquerda,
        legível mesmo com nomes longos; extras garantidos mesmo quando o
        valor do combustível é 0 (ex.: Galão de ferramentas)."""
        por_veic = {}
        ordem = []
        for it in itens:
            vid = it["vid"]
            if vid not in por_veic:
                por_veic[vid] = {"rotulo": it["veiculo_rotulo"], "itens": []}
                ordem.append(vid)
            por_veic[vid]["itens"].append(it)
        ordem = [vid for vid in ordem
                 if any((x["combustivel"] + x["extras"]) > 0
                        for x in por_veic[vid]["itens"])]
        if not ordem:
            return "", ""

        def _fmt(v):
            return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        def _esc(t):
            return (str(t).replace("&", "&amp;").replace("<", "&lt;")
                    .replace(">", "&gt;").replace('"', "&quot;"))

        def _cor(comb):
            return _COMB_CORES.get(comb, "#78909C")

        def _rotulo(t, mx=24):
            t = str(t)
            return t if len(t) <= mx else t[:mx - 1] + "…"

        maxv = max(
            sum(x["combustivel"] + x["extras"] for x in por_veic[vid]["itens"])
            for vid in ordem) or 0.01
        m_e, m_d, m_t, m_b = 160, 70, 18, 10
        alt_linha = 30
        altura = m_t + m_b + alt_linha * len(ordem)
        pw = largura - m_e - m_d
        usados = []
        partes = []
        for grade in range(4):
            gx = m_e + pw * grade / 3
            gv = maxv * grade / 3
            partes.append(f'<line x1="{gx:.1f}" y1="{m_t}" x2="{gx:.1f}" '
                          f'y2="{altura - m_b}" stroke="#E7EDF4"/>')
            partes.append(f'<text x="{gx:.1f}" y="{m_t - 6}" font-size="9" '
                          f'fill="#5A6B7C" text-anchor="middle">'
                          f'R$ {_fmt(gv)}</text>')
        for i, vid in enumerate(ordem):
            info = por_veic[vid]
            y = m_t + i * alt_linha
            alt_barra = 16
            yc = y + (alt_linha - alt_barra) / 2
            total = sum(x["combustivel"] + x["extras"]
                        for x in info["itens"])
            detalhes = []
            acum = 0.0
            g = [f'<g class="g-mes" data-mes="{_esc(info["rotulo"])}">']
            for it in info["itens"]:
                v_comb = it["combustivel"]
                v_ext = it["extras"]
                if v_comb > 0:
                    detalhes.append(
                        f'{_esc(label_combustivel(it["comb"]))}: '
                        f'R$ {_fmt(v_comb)} · Litros {_fmt(it["litros"])}')
                    if it["comb"] not in usados:
                        usados.append(it["comb"])
                    w = v_comb / maxv * pw
                    g.append(f'<rect x="{m_e + acum:.1f}" y="{yc:.1f}" '
                             f'width="{max(w, 1):.1f}" height="{alt_barra}" '
                             f'fill="{_cor(it["comb"])}"/>')
                    acum += w
                if v_ext > 0:
                    detalhes.append(f'Itens extras: R$ {_fmt(v_ext)}')
                    if "extras" not in usados:
                        usados.append("extras")
                    w = v_ext / maxv * pw
                    g.append(f'<rect x="{m_e + acum:.1f}" y="{yc:.1f}" '
                             f'width="{max(w, 1):.1f}" height="{alt_barra}" '
                             f'fill="#E6A23C" opacity=".85"/>')
                    acum += w
            g.append(f'<title>{_esc(info["rotulo"])} — ' +
                     ("; ".join(detalhes) if detalhes else "Sem custos") +
                     f' · Total R$ {_fmt(total)}</title>')
            g.append(f'<rect class="captura" data-mes="{_esc(info["rotulo"])}" '
                     f'data-det="{"||".join(detalhes) if detalhes else "Sem custos."}||Total: R$ {_fmt(total)}" '
                     f'x="{m_e}" y="{y}" width="{pw}" height="{alt_linha}" '
                     f'fill="transparent"/>')
            g.append('</g>')
            partes.append("".join(g))
            partes.append(
                f'<text x="{m_e - 8}" y="{y + alt_linha / 2 + 3:.1f}" '
                f'font-size="10" fill="#22303F" text-anchor="end">'
                f'{_esc(_rotulo(info["rotulo"]))}</text>')
            partes.append(
                f'<text x="{m_e + acum + 5:.1f}" '
                f'y="{y + alt_linha / 2 + 3:.1f}" font-size="9" '
                f'fill="#43505D">R$ {_fmt(total)}</text>')
        rot_leg = [("extras", "Itens extras")] + \
                  [(c, label_combustivel(c)) for c in usados if c != "extras"]
        leg = "".join(
            f'<span class="leg-item"><span class="leg-cor" '
            f'style="background:{("#E6A23C" if c == "extras" else _cor(c))}">'
            f'</span>{_esc(rot)}</span>' for c, rot in rot_leg)
        return Markup(
            f'<svg viewBox="0 0 {largura} {altura}" width="{largura}" '
            f'height="{altura}" role="img" '
            f'style="width:100%;height:auto;" '
            f'xmlns="http://www.w3.org/2000/svg">{"".join(partes)}</svg>'), leg

    def _funcionarios_nomes() -> list:
        try:
            from src.core.employee_repo import EmployeeRepository
            return sorted(
                (e.nome or "" for e in
                 EmployeeRepository().get_all(limit=100000) if e.nome))
        except Exception:
            return []

    # ================= veiculos =================

    _STATUS_FROTA = {
        "viagem": ("Em Viagem", "b-azul"),
        "manutencao": ("Em Manutenção", "b-amarelo"),
        "indisponivel": ("Indisponível", "b-cinza"),
        "disponivel": ("Disponível", "b-verde"),
    }

    def _status_veiculo(repo, v: dict):
        """Deriva status/motorista/destino: viagem > manutenção/indisponível > disponível."""
        vid = v["id"]
        movs = repo.list_movimentacoes(vid)
        aberta = next((m for m in movs if m["aberta"]), None)
        pior = None
        for m in repo.list_manutencoes(vid):
            if pior is None or m["restante"] < pior["restante"]:
                pior = m
        if aberta:
            status, motorista, destino = "viagem", aberta["motorista"], aberta["destino"]
        elif pior and pior["status"] == "vencido":
            status, motorista, destino = "indisponivel", "", ""
        elif pior and pior["status"] == "urgente":
            status, motorista, destino = "manutencao", "", ""
        else:
            status, motorista, destino = "disponivel", "", ""
        if status != "viagem" and not v["proprio"] and v.get("fim_contrato_aluguel") \
                and v["fim_contrato_aluguel"] < date.today().isoformat():
            status = "indisponivel"
        label, classe = _STATUS_FROTA[status]
        v["status"] = status
        v["status_label"] = label
        v["status_classe"] = classe
        v["motorista_atual"] = motorista or ""
        v["destino_atual"] = destino or ""

    @app.get("/frota")
    def frota_lista(request: Request, busca: str = "", page: int = 1,
                    per: int = 20,
                    user: dict = auth.require_permission("frota")):
        repo = _repo()
        per_val = per if per in _PER_OPCOES else 20
        pagina = max(1, page)
        veiculos, total = repo.list_veiculos(busca=busca,
                                             limit=per_val,
                                             offset=(pagina - 1) * per_val)
        for v in veiculos:
            v["rotulo"] = veiculo_rotulo(v)
            v["tipo_label"] = label_tipo(v["tipo"])
            v["sub_label"] = label_subtipo(v["subtipo"])
            v["posse"] = "Próprio" if v["proprio"] else \
                f"Alugado — {v['contratante'] or '?'}"
            _status_veiculo(repo, v)
        page_n, paginas = _paginacao(request, total, pagina, per_val)
        qs = f"busca={quote_plus((busca or '').strip())}" \
            if (busca or "").strip() else ""
        if per_val != 20:
            qs = (qs + "&" if qs else "") + f"per={per_val}"
        return templates.TemplateResponse(
            request=request, name="frota.html",
            context=ctx(request, veiculos=veiculos, total=total,
                        busca=(busca or "").strip(), page=page_n,
                        paginas=paginas,
                        per=per_val, per_opcoes=_PER_OPCOES,
                        pg_base=("/frota?" + qs) if qs else "/frota",
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    def _form_context(v=None, erro: str = "", valores: dict = None):
        emp_opts = sorted(_repo().list_empresas(), key=lambda e: e["nome"])
        vals = valores or {}
        return dict(
            v=v, erro=erro, emp_opts=emp_opts,
            tipos=TIPOS_VEICULO, subtipos=SUBTIPOS_CAMINHAO,
            carroc=CARROCERIAS,
            f_modelo=vals.get("modelo", v["modelo"] if v else ""),
            f_marca=vals.get("marca", v["marca"] if v else ""),
            f_tipo=vals.get("tipo", v["tipo"] if v else "carro"),
            f_subtipo=vals.get("subtipo", v["subtipo"] if v else ""),
            f_placa=vals.get("placa", v["placa"] if v else ""),
            f_proprio=vals.get("proprio", bool(v["proprio"]) if v else True),
            f_contratante=vals.get("contratante",
                                   v["contratante"] if v else ""),
            f_empresa_id=vals.get("empresa_id",
                                  v["empresa_id"] if v else None),
            f_obs=vals.get("obs", v["obs"] if v else ""),
            f_cor=vals.get("cor", v["cor"] if v else ""),
            f_carroceria=vals.get("carroceria",
                                  v["carroceria"] if v else ""),
            f_ano=vals.get("ano", v["ano"] if v else ""),
            f_fim_contrato=vals.get("fim_contrato_aluguel",
                                    _br(v["fim_contrato_aluguel"])
                                    if v else ""),
            f_km_l=vals.get("km_l_esperado",
                            v["km_l_esperado"] if v else ""),
        )

    @app.get("/frota/novo")
    def frota_novo(request: Request,
                   user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota")
        if red:
            return red
        return templates.TemplateResponse(
            request=request, name="frota_form.html",
            context=ctx(request, **_form_context()))

    def _campos_extras(cor: str, carroceria: str, ano: str,
                       fim_contrato: str, km_l: str):
        fim_iso = _iso(fim_contrato)
        if fim_contrato and fim_contrato.strip() and fim_iso is None:
            raise ValueError("Data de fim do contrato inválida (dd/mm/aaaa).")
        try:
            km_l_val = float(str(km_l).replace(",", "."))
        except (TypeError, ValueError):
            km_l_val = None
        return dict(cor=cor, carroceria=carroceria, ano=ano,
                    fim_contrato_aluguel=fim_iso, km_l_esperado=km_l_val)

    @app.post("/frota/criar")
    def frota_criar(request: Request, modelo: str = Form(""),
                    marca: str = Form(""), tipo: str = Form(""),
                    subtipo: str = Form(""), placa: str = Form(""),
                    proprio: str = Form(""), contratante: str = Form(""),
                    empresa_id: str = Form(""), obs: str = Form(""),
                    cor: str = Form(""), carroceria: str = Form(""),
                    ano: str = Form(""), fim_contrato: str = Form(""),
                    km_l: str = Form(""),
                    user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota")
        if red:
            return red
        vals = dict(modelo=modelo, marca=marca, tipo=tipo, subtipo=subtipo,
                    placa=placa, proprio=proprio == "1",
                    contratante=contratante, empresa_id=empresa_id, obs=obs,
                    cor=cor, carroceria=carroceria, ano=ano,
                    fim_contrato_aluguel=fim_contrato, km_l_esperado=km_l)
        try:
            extras = _campos_extras(cor, carroceria, ano, fim_contrato, km_l)
            emp_id = int(empresa_id) if empresa_id else None
            novo = _repo().add_veiculo(
                modelo, marca, tipo, subtipo, placa, vals["proprio"],
                contratante, emp_id, obs, **extras)
        except ValueError as e:
            return templates.TemplateResponse(
                request=request, name="frota_form.html",
                context=ctx(request, **_form_context(erro=str(e),
                                                     valores=vals)))
        _audit(request, "frota-veiculo-criar", user["username"],
               str(novo), f"{modelo} {placa}")
        flash(request, msg="Veículo cadastrado.")
        return RedirectResponse(f"/frota/{novo}", status_code=303)

    @app.get("/frota/{veiculo_id}/editar")
    def frota_editar(veiculo_id: int, request: Request,
                     user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        v = _repo().get_veiculo(veiculo_id)
        if not v:
            flash(request, erro="Veículo não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        return templates.TemplateResponse(
            request=request, name="frota_form.html",
            context=ctx(request, **_form_context(v=v)))

    @app.post("/frota/{veiculo_id}/editar")
    def frota_salvar(veiculo_id: int, request: Request,
                     modelo: str = Form(""), marca: str = Form(""),
                     tipo: str = Form(""), subtipo: str = Form(""),
                     placa: str = Form(""), proprio: str = Form(""),
                     contratante: str = Form(""), empresa_id: str = Form(""),
                     obs: str = Form(""), cor: str = Form(""),
                     carroceria: str = Form(""), ano: str = Form(""),
                     fim_contrato: str = Form(""), km_l: str = Form(""),
                     user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        vals = dict(modelo=modelo, marca=marca, tipo=tipo, subtipo=subtipo,
                    placa=placa, proprio=proprio == "1",
                    contratante=contratante, empresa_id=empresa_id, obs=obs,
                    cor=cor, carroceria=carroceria, ano=ano,
                    fim_contrato_aluguel=fim_contrato, km_l_esperado=km_l)
        try:
            extras = _campos_extras(cor, carroceria, ano, fim_contrato, km_l)
            emp_id = int(empresa_id) if empresa_id else None
            _repo().update_veiculo(
                veiculo_id, modelo, marca, tipo, subtipo, placa,
                vals["proprio"], contratante, emp_id, obs, **extras)
        except ValueError as e:
            v = _repo().get_veiculo(veiculo_id)
            return templates.TemplateResponse(
                request=request, name="frota_form.html",
                context=ctx(request, **_form_context(v=v, erro=str(e),
                                                     valores=vals)))
        _audit(request, "frota-veiculo-editar", user["username"],
               str(veiculo_id), f"{modelo} {placa}")
        flash(request, msg="Veículo atualizado.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.post("/frota/{veiculo_id}/excluir")
    def frota_excluir(veiculo_id: int, request: Request,
                      user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota")
        if red:
            return red
        v = _repo().get_veiculo(veiculo_id)
        if not v:
            flash(request, erro="Veículo não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        try:
            _repo().delete_veiculo(veiculo_id)
            _audit(request, "frota-veiculo-excluir", user["username"],
                   str(veiculo_id), veiculo_rotulo(v))
            flash(request, msg=f"Veículo '{veiculo_rotulo(v)}' excluído.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota", status_code=303)

    @app.post("/frota/{veiculo_id}/foto")
    async def frota_foto_upload(veiculo_id: int, request: Request,
                                arquivo: UploadFile = File(None),
                                user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione uma imagem.")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        data = await arquivo.read()
        ext = Path(arquivo.filename).suffix.lower().lstrip(".")
        tipo = {"jpg": "jpg", "jpeg": "jpg", "png": "png", "gif": "gif",
                "webp": "webp"}.get(ext)
        try:
            if tipo is None:
                raise ValueError("Formato não suportado (use JPG, PNG, GIF"
                                 " ou WEBP).")
            _repo().update_foto(veiculo_id, data, tipo)
        except ValueError as e:
            flash(request, erro=str(e))
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        flash(request, msg="Foto atualizada.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.get("/frota/{veiculo_id}/foto")
    def frota_foto(veiculo_id: int):
        f = _repo().get_foto(veiculo_id)
        if not f:
            return Response(status_code=404)
        return Response(content=f["dados"],
                        media_type=f"image/{'jpeg' if f['tipo'] == 'jpg' else f['tipo']}")

    @app.get("/frota/exportar")
    def frota_exportar(request: Request,
                       user: dict = auth.require_permission("frota")):
        """2.29.7: exporta a frota para Excel (base para reimportar)."""
        import io
        import openpyxl
        veiculos, _t = _repo().list_veiculos(limit=100000, offset=0)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Veiculos"
        ws.append(["Modelo", "Marca", "Tipo", "Subtipo", "Placa", "Proprio",
                   "Contratante", "Empresa", "Cor", "Carroceria", "Ano",
                   "Fim do contrato", "KM/L esperado", "Obs"])
        for v in veiculos:
            ws.append([
                v["modelo"], v["marca"] or "", label_tipo(v["tipo"]),
                label_subtipo(v["subtipo"]) if v["subtipo"] else "",
                v["placa"] or "", "Sim" if v["proprio"] else "Não",
                v["contratante"] or "", v["empresa_nome"] or "",
                v["cor"] or "", label_subtipo(v["carroceria"])
                if v["carroceria"] in ("hatch", "sedan") else (v["carroceria"]
                                                               or ""),
                v["ano"] or "",
                _br(v["fim_contrato_aluguel"])
                if v["fim_contrato_aluguel"] else "",
                v["km_l_esperado"] or "", v["obs"] or "",
            ])
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        _audit(request, "frota-exportar", user["username"], "",
               f"{len(veiculos)} veiculo(s)")
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument"
                       ".spreadsheetml.sheet",
            headers={"Content-Disposition":
                     'attachment; filename="veiculos.xlsx"'})

    @app.post("/frota/importar")
    async def frota_importar(request: Request,
                             arquivo: UploadFile = File(None),
                             user: dict = auth.require_permission("frota")):
        """2.29.7: importa veículos de planilha (modelo MODELO VEICULOS.xlsx)."""
        red = _bloqueio(request, user, "/frota")
        if red:
            return red
        from src.web.routers.importacoes import ler_upload_xlsx
        caminho, erro = await ler_upload_xlsx(arquivo)
        if erro:
            flash(request, erro=erro)
            return RedirectResponse("/frota", status_code=303)
        try:
            from src.utils.veiculo_importer import import_veiculos_from_excel
            importados, erros = import_veiculos_from_excel(
                caminho, _repo())
        except Exception as e:
            from src.utils.error_log import log_error
            log_error("portal-frota-importar", e)
            flash(request, erro=f"Falha ao ler a planilha: {e}")
            return RedirectResponse("/frota", status_code=303)
        finally:
            try:
                caminho.unlink()
            except OSError:
                pass
        if erros:
            resumo = "; ".join(erros[:5])
            if len(erros) > 5:
                resumo += f" (e mais {len(erros) - 5} erro(s))"
            flash(request, msg=f"{importados} veículo(s) importado(s).",
                  erro=resumo)
        else:
            flash(request, msg=f"{importados} veículo(s) importado(s).")
        _audit(request, "frota-importar", user["username"], "",
               f"{importados} importado(s), {len(erros)} erro(s)")
        return RedirectResponse("/frota", status_code=303)

    # ================= empresas de veiculos =================

    @app.get("/frota/empresas")
    def frota_empresas(request: Request,
                       user: dict = auth.require_permission("frota")):
        return templates.TemplateResponse(
            request=request, name="frota_empresas.html",
            context=ctx(request, empresas=_repo().list_empresas(),
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.post("/frota/empresas/criar")
    def frota_empresa_criar(request: Request, nome: str = Form(""),
                            cnpj: str = Form(""),
                            user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/empresas")
        if red:
            return red
        erro_cnpj = _cnpj_invalido(cnpj)
        if erro_cnpj:
            flash(request, erro=erro_cnpj)
            return RedirectResponse("/frota/empresas", status_code=303)
        try:
            _repo().add_empresa(nome, cnpj)
            flash(request, msg="Empresa cadastrada.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/empresas", status_code=303)

    @app.post("/frota/empresas/{empresa_id}/editar")
    def frota_empresa_editar(empresa_id: int, request: Request,
                             nome: str = Form(""), cnpj: str = Form(""),
                             user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/empresas")
        if red:
            return red
        erro_cnpj = _cnpj_invalido(cnpj)
        if erro_cnpj:
            flash(request, erro=erro_cnpj)
            return RedirectResponse("/frota/empresas", status_code=303)
        try:
            _repo().update_empresa(empresa_id, nome, cnpj)
            flash(request, msg="Empresa atualizada.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/empresas", status_code=303)

    @app.post("/frota/empresas/{empresa_id}/excluir")
    def frota_empresa_excluir(empresa_id: int, request: Request,
                              user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/empresas")
        if red:
            return red
        try:
            _repo().delete_empresa(empresa_id)
            flash(request, msg="Empresa excluída.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/empresas", status_code=303)

    # ================= fornecedores =================

    @app.get("/frota/fornecedores")
    def frota_fornecedores(request: Request,
                           user: dict = auth.require_permission("frota")):
        return templates.TemplateResponse(
            request=request, name="frota_fornecedores.html",
            context=ctx(request, fornecedores=_repo().list_fornecedores(),
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.post("/frota/fornecedores/criar")
    def frota_fornecedor_criar(request: Request, nome: str = Form(""),
                               cnpj: str = Form(""),
                               endereco: str = Form(""),
                               user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/fornecedores")
        if red:
            return red
        erro_cnpj = _cnpj_invalido(cnpj)
        if erro_cnpj:
            flash(request, erro=erro_cnpj)
            return RedirectResponse("/frota/fornecedores", status_code=303)
        try:
            _repo().add_fornecedor(nome, cnpj, endereco)
            flash(request, msg="Fornecedor cadastrado.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/fornecedores", status_code=303)

    @app.post("/frota/fornecedores/{forn_id}/editar")
    def frota_fornecedor_editar(forn_id: int, request: Request,
                                nome: str = Form(""), cnpj: str = Form(""),
                                endereco: str = Form(""),
                                user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/fornecedores")
        if red:
            return red
        erro_cnpj = _cnpj_invalido(cnpj)
        if erro_cnpj:
            flash(request, erro=erro_cnpj)
            return RedirectResponse("/frota/fornecedores", status_code=303)
        try:
            _repo().update_fornecedor(forn_id, nome, cnpj, endereco)
            flash(request, msg="Fornecedor atualizado.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/fornecedores", status_code=303)

    @app.post("/frota/fornecedores/{forn_id}/excluir")
    def frota_fornecedor_excluir(forn_id: int, request: Request,
                                 user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/fornecedores")
        if red:
            return red
        try:
            _repo().delete_fornecedor(forn_id)
            flash(request, msg="Fornecedor excluído.")
        except ValueError as e:
            flash(request, erro=str(e))
        return RedirectResponse("/frota/fornecedores", status_code=303)

    # ================= abastecimentos =================

    @app.get("/frota/abastecimentos")
    def frota_abast_lista(request: Request, busca: str = "", page: int = 1,
                          per: int = 20, ordem: str = "data",
                          situacao: str = "todas",
                          user: dict = auth.require_permission("frota")):
        repo = _repo()
        per_val = per if per in _PER_OPCOES else 20
        ordem_val = ordem if ordem in ("data", "serial") else "data"
        sit_val = (situacao if situacao in ("todas", "ativas", "bloqueadas")
                   else "todas")
        pagina = max(1, page)
        itens, total = repo.list_abastecimentos(
            busca=busca, limit=per_val, offset=(pagina - 1) * per_val,
            ordem=ordem_val, situacao=sit_val)
        nf_map = repo.tem_nf([a["id"] for a in itens])
        for a in itens:
            a["data_br"] = _br(a["data"])
            nf_num = nf_map.get(a["id"])
            a["nf_numero"] = nf_num
            a["tem_nf"] = bool(nf_num)
            a["bloqueada"] = (a.get("status") == "bloqueada")
            a["revisao_rotulo"] = rotulo_revisao(a.get("revisao"))
            a["pode_excluir"] = repo.pode_excluir_abastecimento(a["id"])
        page_n, paginas = _paginacao(request, total, pagina, per_val)
        qs = []
        if (busca or "").strip():
            qs.append(f"busca={quote_plus(busca.strip())}")
        if per_val != 20:
            qs.append(f"per={per_val}")
        if ordem_val != "data":
            qs.append(f"ordem={ordem_val}")
        if sit_val != "todas":
            qs.append(f"situacao={sit_val}")
        pg = ("/frota/abastecimentos?" + "&".join(qs)) if qs \
            else "/frota/abastecimentos"
        return templates.TemplateResponse(
            request=request, name="frota_abastecimentos.html",
            context=ctx(request, itens=itens, total=total,
                        busca=(busca or "").strip(), page=page_n,
                        paginas=paginas,
                        per=per_val, per_opcoes=_PER_OPCOES,
                        ordem=ordem_val, situacao=sit_val,
                        motivos_bloqueio=_MOTIVOS_BLOQUEIO,
                        pg_base=pg,
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.get("/frota/abastecimentos/novo")
    def frota_abast_novo(request: Request, veiculo: int = 0,
                         user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        repo = _repo()
        veiculo_sel = veiculo if repo.get_veiculo(veiculo) else None
        return templates.TemplateResponse(
            request=request, name="frota_abast_form.html",
            context=ctx(request, erro="",
                        veiculo_sel=veiculo_sel,
                        funcionarios=_funcionarios_nomes(),
                        veiculos=sorted(repo.list_veiculos(limit=500)[0],
                                        key=lambda v: (v["modelo"] or "",
                                                       v["marca"] or "")),
                        fornecedores=sorted(repo.list_fornecedores(),
                                            key=lambda f: f["nome"]),
                        combustiveis=TIPOS_COMBUSTIVEL,
                        hoje=date.today().strftime("%d/%m/%Y")))

    @app.post("/frota/abastecimentos/criar")
    async def frota_abast_criar(request: Request, veiculo_id: str = Form(""),
                                fornecedor_id: str = Form(""),
                                combustivel: str = Form(""), data: str = Form(""),
                                viagem_servico: str = Form(""), km: str = Form(""),
                                condutor: str = Form(""),
                                obs: str = Form(""), litros: str = Form(""),
                                valor: str = Form(""),
                                usar_tmp: str = Form(""),
                                tmp_fornecedor: str = Form(""),
                                tmp_cnpj: str = Form(""),
                                user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        extras = _ler_extras(await request.form())
        data_iso = _iso(data)
        repo = _repo()

        def _re_render(erro):
            return templates.TemplateResponse(
                request=request, name="frota_abast_form.html",
                context=ctx(request, erro=erro,
                            veiculo_sel=None,
                            funcionarios=_funcionarios_nomes(),
                            veiculos=sorted(
                                repo.list_veiculos(limit=500)[0],
                                key=lambda v: (v["modelo"] or "",
                                               v["marca"] or "")),
                            fornecedores=sorted(repo.list_fornecedores(),
                                                key=lambda f: f["nome"]),
                            combustiveis=TIPOS_COMBUSTIVEL,
                            hoje=data.strip() or
                            date.today().strftime("%d/%m/%Y")))

        if data_iso is None:
            return _re_render("Data inválida (dd/mm/aaaa).")
        try:
            def _num(txt):
                return float(str(txt).replace(",", ".")) if str(txt).strip() \
                    else None
            vid = int(veiculo_id)
            # 2.29.7: combustível compatível com o tipo do veículo
            v_info = repo.get_veiculo(vid)
            if not v_info:
                return _re_render("Veículo inválido.")
            if v_info["tipo"] in _TIPOS_SEM_DIESEL and \
                    combustivel in _COMB_BLOQUEADOS_LEVES:
                from src.core.frota_repo import label_combustivel as _lc
                return _re_render(
                    f"{_lc(combustivel)} não se aplica a "
                    f"{label_tipo(v_info['tipo'])}.")
            fid = int(fornecedor_id) if fornecedor_id else None
            km_val = int(km) if km.strip() else None
            tmp_f = tmp_c = ""
            if usar_tmp == "1":
                tmp_f = (tmp_fornecedor or "").strip()
                tmp_c = (tmp_cnpj or "").strip()
                if not tmp_f:
                    return _re_render(
                        "Informe o nome do posto temporário (ou desmarque "
                        "a opção).")
                err_cnpj = _cnpj_invalido(tmp_c) if tmp_c else None
                if err_cnpj:
                    return _re_render(f"Posto temporário: {err_cnpj}")
                fid = None
            abast_id, serial = repo.add_abastecimento(
                vid, fid, combustivel, data_iso, viagem_servico, km_val,
                condutor, obs=obs, extras=extras, litros=_num(litros),
                valor=_num(valor), tmp_fornecedor=tmp_f, tmp_cnpj=tmp_c)
        except (ValueError, TypeError) as e:
            return _re_render(str(e))
        # PDF
        pdf_path = None
        try:
            from src.core.pdf_abastecimento import gerar_pdf_abastecimento
            from src.core.config import load_company_config
            abast = repo.get_abastecimento(abast_id)
            pdf = gerar_pdf_abastecimento({
                "serial": serial,
                "data_br": _br(abast["data"]),
                "veiculo": abast,
                "fornecedor": {"nome": abast.get("fornecedor"),
                               "cnpj": abast.get("fornecedor_cnpj"),
                               "endereco": abast.get("fornecedor_endereco")},
                "combustivel": abast["combustivel"],
                "condutor": abast["condutor"],
                "viagem_servico": abast["viagem_servico"],
                "km": abast["km"],
                "obs": abast["obs"],
                "extras": abast.get("extras_lista") or [],
                "extras_total": abast.get("extras_total"),
                "config": load_company_config(),
            })
            pdf_path = str(pdf)
            repo.set_pdf_path(abast_id, pdf_path)
        except Exception as e:
            from src.utils.error_log import log_error
            log_error("portal-frota-pdf", e)
        _audit(request, "frota-abastecimento", user["username"],
               serial, f"veiculo={vid}")
        flash(request, msg=f"Solicitação {serial} registrada.")
        return RedirectResponse("/frota/abastecimentos", status_code=303)

    @app.get("/frota/abastecimentos/{abast_id}/editar")
    def frota_abast_editar(abast_id: int, request: Request,
                           user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        return templates.TemplateResponse(
            request=request, name="frota_abast_form.html",
            context=ctx(request, erro="",
                        abast=a,
                        data_br=_br(a["data"]),
                        veiculo_rotulo=veiculo_rotulo(a),
                        funcionarios=_funcionarios_nomes(),
                        fornecedores=sorted(repo.list_fornecedores(),
                                            key=lambda f: f["nome"]),
                        combustiveis=TIPOS_COMBUSTIVEL,
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.post("/frota/abastecimentos/{abast_id}/editar")
    async def frota_abast_salvar(abast_id: int, request: Request,
                                 fornecedor_id: str = Form(""),
                                 combustivel: str = Form(""), data: str = Form(""),
                                 viagem_servico: str = Form(""), km: str = Form(""),
                                 condutor: str = Form(""),
                                 obs: str = Form(""), litros: str = Form(""),
                                 valor: str = Form(""),
                                 usar_tmp: str = Form(""),
                                 tmp_fornecedor: str = Form(""),
                                 tmp_cnpj: str = Form(""),
                                 user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        extras = _ler_extras(await request.form())
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        data_iso = _iso(data)

        def _re_render(erro):
            return templates.TemplateResponse(
                request=request, name="frota_abast_form.html",
                context=ctx(request, erro=erro,
                            abast=a,
                            data_br=data.strip() or _br(a["data"]),
                            veiculo_rotulo=veiculo_rotulo(a),
                            funcionarios=_funcionarios_nomes(),
                            fornecedores=sorted(
                                repo.list_fornecedores(),
                                key=lambda f: f["nome"]),
                            combustiveis=TIPOS_COMBUSTIVEL,
                            pode_escrever=auth.pode_escrever(
                                user["papel"], "frota")))

        if data_iso is None:
            return _re_render("Data inválida (dd/mm/aaaa).")
        try:
            v_info = repo.get_veiculo(a["veiculo_id"])
            if v_info and v_info["tipo"] in _TIPOS_SEM_DIESEL and \
                    combustivel in _COMB_BLOQUEADOS_LEVES:
                return _re_render(
                    f"{label_combustivel(combustivel)} não se aplica a "
                    f"{label_tipo(v_info['tipo'])}.")
            fid = int(fornecedor_id) if fornecedor_id else None
            km_val = int(km) if km.strip() else None
            tmp_f = tmp_c = ""
            if usar_tmp == "1":
                tmp_f = (tmp_fornecedor or "").strip()
                tmp_c = (tmp_cnpj or "").strip()
                if not tmp_f:
                    return _re_render(
                        "Informe o nome do posto temporário (ou desmarque "
                        "a opção).")
                err_cnpj = _cnpj_invalido(tmp_c) if tmp_c else None
                if err_cnpj:
                    return _re_render(f"Posto temporário: {err_cnpj}")
                fid = None
            litros_v = float(str(litros).replace(",", ".")) \
                if str(litros).strip() else None
            valor_v = float(str(valor).replace(",", ".")) \
                if str(valor).strip() else None
            rev = repo.update_abastecimento(
                abast_id, fid, combustivel, data_iso, viagem_servico, km_val,
                condutor, obs, litros=litros_v, valor=valor_v, extras=extras,
                tmp_fornecedor=tmp_f, tmp_cnpj=tmp_c)
        except (ValueError, TypeError) as e:
            return _re_render(str(e))
        # PDF regenerado com a marca de revisão (2.33.2)
        try:
            from src.core.pdf_abastecimento import gerar_pdf_abastecimento
            from src.core.config import load_company_config
            abast = repo.get_abastecimento(abast_id)
            pdf = gerar_pdf_abastecimento({
                "serial": abast["serial"],
                "data_br": _br(abast["data"]),
                "veiculo": abast,
                "fornecedor": {"nome": abast.get("fornecedor"),
                               "cnpj": abast.get("fornecedor_cnpj"),
                               "endereco": abast.get("fornecedor_endereco")},
                "combustivel": abast["combustivel"],
                "condutor": abast["condutor"],
                "viagem_servico": abast["viagem_servico"],
                "km": abast["km"],
                "obs": abast["obs"],
                "extras": abast.get("extras_lista") or [],
                "extras_total": abast.get("extras_total"),
                "revisao": rotulo_revisao(rev),
                "config": load_company_config(),
            })
            repo.set_pdf_path(abast_id, str(pdf))
        except Exception as e:
            from src.utils.error_log import log_error
            log_error("portal-frota-pdf-rev", e)
        _audit(request, "frota-abastecimento-editar", user["username"],
               a["serial"], f"revisao={rev}")
        flash(request, msg=(f"Solicitação {a['serial']} atualizada "
                            f"({rotulo_revisao(rev)})."))
        return RedirectResponse("/frota/abastecimentos", status_code=303)

    # ---------- bloqueio / exclusão (2.35.2) ----------

    @app.post("/frota/abastecimentos/{abast_id}/bloquear")
    async def frota_abast_bloquear(abast_id: int, request: Request,
                                   user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        fdata = await request.form()
        motivo_sel = str(fdata.get("motivo") or "").strip()
        motivo_txt = str(fdata.get("motivo_txt") or "").strip()
        if motivo_sel == "Outro" and motivo_txt:
            motivo = motivo_txt
        else:
            motivo = motivo_sel if motivo_sel in _MOTIVOS_BLOQUEIO \
                else "Outro"
        if repo.bloquear_abastecimento(abast_id, motivo, user["username"]):
            _audit(request, "frota-abastecimento-bloquear", user["username"],
                   a["serial"], motivo)
            flash(request, msg=(f"Solicitação {a['serial']} bloqueada "
                                f"({motivo}) — saiu dos totais de custo."))
        else:
            flash(request, erro="Não foi possível bloquear a solicitação.")
        return RedirectResponse("/frota/abastecimentos", status_code=303)

    @app.post("/frota/abastecimentos/{abast_id}/desbloquear")
    def frota_abast_desbloquear(abast_id: int, request: Request,
                                user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        repo.desbloquear_abastecimento(abast_id)
        _audit(request, "frota-abastecimento-desbloquear", user["username"],
               a["serial"], "")
        flash(request, msg=f"Solicitação {a['serial']} reativada.")
        return RedirectResponse("/frota/abastecimentos", status_code=303)

    @app.post("/frota/abastecimentos/{abast_id}/excluir")
    def frota_abast_excluir(abast_id: int, request: Request,
                            user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, "/frota/abastecimentos")
        if red:
            return red
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        if not repo.pode_excluir_abastecimento(abast_id):
            flash(request, erro=("Somente a solicitação mais recente pode ser "
                                 "excluída. Bloqueie esta para tirá-la dos custos."))
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        if repo.delete_abastecimento(abast_id):
            try:
                if a.get("pdf_path"):
                    Path(a["pdf_path"]).unlink(missing_ok=True)
            except Exception:
                pass
            _audit(request, "frota-abastecimento-excluir", user["username"],
                   a["serial"], "")
            flash(request, msg=f"Solicitação {a['serial']} excluída.")
        else:
            flash(request, erro="Não foi possível excluir a solicitação.")
        return RedirectResponse("/frota/abastecimentos", status_code=303)

    @app.get("/frota/abastecimentos/exportar")
    def frota_abast_exportar(request: Request,
                             user: dict = auth.require_permission("frota")):
        import io
        import openpyxl
        repo = _repo()
        itens, _total = repo.list_abastecimentos(limit=5000, offset=0)
        nf_map = repo.tem_nf([a["id"] for a in itens])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Abastecimentos"
        ws.append(["Serial", "Data", "Veículo", "Combustível", "KM", "Litros",
                   "Valor (R$)", "Itens extras (R$)", "Total (R$)",
                   "NF", "Situação", "Motivo", "Fornecedor", "Condutor",
                   "Viagem/Serviço", "Observações"])
        for a in itens:
            nf_num = nf_map.get(a["id"])
            extras = float(a.get("extras_total") or 0)
            valor = float(a.get("valor") or 0)
            bloqueada = (a.get("status") == "bloqueada")
            ws.append([
                a["serial"], _br(a["data"]), a["veiculo_rotulo"],
                a["combustivel_label"], a["km"], a["litros"], valor,
                extras if extras else None,
                (valor + extras) if (valor or extras) else None,
                nf_num or "Sem NF",
                "Bloqueada" if bloqueada else "Ativa",
                a.get("motivo_status") or "",
                a.get("fornecedor") or "", a.get("condutor") or "",
                a.get("viagem_servico") or "", a.get("obs") or "",
            ])
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=9):
            for cell in row:
                cell.number_format = "0.00"
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument"
                       ".spreadsheetml.sheet",
            headers={"Content-Disposition":
                     'attachment; filename="abastecimentos.xlsx"'})

    # ---------- exportação de custos (2.35.1) ----------

    _XLSX_MEDIA = ("application/vnd.openxmlformats-officedocument"
                   ".spreadsheetml.sheet")

    def _abast_sheet(ws, linhas, nf_map) -> None:
        ws.append(["Serial", "Data", "Veículo", "Combustível", "KM", "Litros",
                   "Combustível (R$)", "Extras (R$)", "Total (R$)",
                   "NF", "Situação", "Motivo"])
        for a in linhas:
            nf_num = (nf_map or {}).get(a["id"])
            extras = float(a.get("extras_total") or 0)
            valor = float(a.get("valor") or 0)
            bloqueada = (a.get("status") == "bloqueada")
            ws.append([
                a["serial"], _br(a["data"]),
                a.get("veiculo_rotulo") or "", a.get("combustivel_label") or "",
                a.get("km"), a.get("litros"), valor,
                extras if extras else None,
                (valor + extras) if (valor or extras) else None,
                nf_num or "Sem NF",
                "Bloqueada" if bloqueada else "Ativa",
                a.get("motivo_status") or "",
            ])
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=9):
            for cell in row:
                cell.number_format = "0.00"

    @app.get("/frota/custos")
    def frota_custos(request: Request,
                     user: dict = auth.require_permission("frota")):
        """Gráficos de custos gerais (2.37.2): 12 meses por veículo e
        por veículo × combustível (incl. extras), com exportações."""
        repo = _repo()
        serie = repo.custo_serie_todos()
        resumos = [r for r in repo.resumo_custo_todos() if r["qtd"]]
        comb = repo.custo_por_combustivel()
        g_meses, leg_meses = _svg_custos_todos(serie)
        g_comb, leg_comb = _svg_custos_comb(comb)
        return templates.TemplateResponse(
            request=request, name="custos_frota.html",
            context=ctx(request, grafico_meses=g_meses, legenda_meses=leg_meses,
                        grafico_comb=g_comb, legenda_comb=leg_comb,
                        resumos=resumos,
                        pode_escrever=pode_escrever(user["papel"], "frota")))

    @app.get("/frota/custos/pdf")
    def frota_custos_pdf(request: Request,
                         user: dict = auth.require_permission("frota")):
        """RELATÓRIO DE CUSTOS — FROTA em PDF (2.37.2), direto no navegador."""
        from src.core.pdf_custos_frota import gerar_pdf_custos_frota
        try:
            data = gerar_pdf_custos_frota()
        except Exception:
            from src.utils.error_log import log_error
            log_error("portal-frota-custos-pdf")
            flash(request, erro="Não foi possível gerar o relatório de custos.")
            return RedirectResponse("/frota/custos", status_code=303)
        _audit(request, "frota-custos-pdf", user["username"], "geral")
        return Response(content=data, media_type="application/pdf",
                        headers={"Content-Disposition":
                                 'inline; filename="custos_frota.pdf"'})

    @app.get("/frota/{veiculo_id}/custos/pdf")
    def frota_custo_veiculo_pdf(veiculo_id: int, request: Request,
                                user: dict = auth.require_permission("frota")):
        """PDF individual de custos do veículo (2.37.2), direto no navegador."""
        from src.core.pdf_custos_frota import gerar_pdf_custos_frota
        v = _repo().get_veiculo(veiculo_id)
        if not v:
            return Response(status_code=404)
        try:
            data = gerar_pdf_custos_frota(veiculo_id=veiculo_id)
        except Exception:
            from src.utils.error_log import log_error
            log_error("portal-frota-custos-pdf")
            flash(request, erro="Não foi possível gerar o relatório de custos.")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        _audit(request, "frota-custos-pdf", user["username"],
               v.get("placa") or str(veiculo_id))
        nome = (v.get("placa") or str(veiculo_id)).replace(" ", "_")
        return Response(content=data, media_type="application/pdf",
                        headers={"Content-Disposition":
                                 f'inline; filename="custos_{nome}.pdf"'})

    @app.get("/frota/custos/exportar")
    def frota_custos_exportar(request: Request,
                              user: dict = auth.require_permission("frota")):
        """Excel geral de custos (2.35.1): resumo por veículo + todas as linhas."""
        import io
        import openpyxl
        repo = _repo()
        resumos = repo.resumo_custo_todos()
        itens, _total = repo.list_abastecimentos(limit=5000, offset=0)
        nf_map = repo.tem_nf([a["id"] for a in itens])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumo por veiculo"
        ws.append(["Veículo", "Abastecimentos (ativos)", "Litros",
                   "Combustível (R$)", "Itens extras (R$)", "Total (R$)"])
        for r in resumos:
            combust = float(r["combustivel"] or 0)
            extras = float(r["extras"] or 0)
            ws.append([r["veiculo_rotulo"], int(r["qtd"] or 0),
                       float(r["litros"] or 0), combust,
                       extras if extras else None,
                       (combust + extras) if (combust or extras) else None])
        ws2 = wb.create_sheet("Abastecimentos")
        _abast_sheet(ws2, itens, nf_map)
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=6):
            for cell in row:
                cell.number_format = "0.00"
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        _audit(request, "frota-custos-exportar", user["username"], "",
               f"{len(resumos)} veiculo(s)")
        return Response(
            content=buf.getvalue(), media_type=_XLSX_MEDIA,
            headers={"Content-Disposition":
                     'attachment; filename="custos_frota.xlsx"'})

    @app.get("/frota/{veiculo_id}/custo/exportar")
    def frota_custo_veiculo_exportar(veiculo_id: int, request: Request,
                                     user: dict = auth.require_permission("frota")):
        """Excel de custo/consumo de um veículo (2.35.1)."""
        import io
        import openpyxl
        repo = _repo()
        v = repo.get_veiculo(veiculo_id)
        if not v:
            return Response(status_code=404)
        resumo = repo.resumo_custo_veiculo(veiculo_id)
        serie = repo.custo_serie_veiculo(veiculo_id)
        abasts = repo.list_abast_por_veiculo(veiculo_id)
        nf_map = repo.tem_nf([a["id"] for a in abasts])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.append(["Custo e consumo", ""])
        ws.append(["Veículo", v["veiculo_rotulo"] if v.get("veiculo_rotulo")
                   else veiculo_rotulo(v)])
        ws.append(["Total abastecido (L)", resumo["litros"]])
        ws.append(["Custo combustível (R$)", resumo["valor_combustivel"]])
        ws.append(["Custo itens extras (R$)", resumo["valor_extras"]])
        ws.append(["Custo total (R$)", resumo["valor"]])
        ws.append(["Média KM/L (real)", resumo["media_km_l"]])
        ws.append(["KM/L esperado", resumo["km_l_esperado"]])
        ws.append(["Custo por km (R$)", resumo["custo_km"]])
        ws.append([])
        ws.append(["Série mensal", "Combustível (R$)", "Extras (R$)",
                   "Litros"])
        for s in serie:
            ws.append([s["rotulo"], s["combustivel"], s["extras"],
                       s["litros"]])
        for rr in range(3, 10):
            ws.cell(row=rr, column=2).number_format = "0.00"
        for rr in range(11, 11 + len(serie)):
            for cc in (2, 3, 4):
                ws.cell(row=rr, column=cc).number_format = "0.00"
        ws2 = wb.create_sheet("Abastecimentos")
        _abast_sheet(ws2, abasts, nf_map)
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        _audit(request, "frota-custo-exportar", user["username"],
               v["veiculo_rotulo"] if v.get("veiculo_rotulo")
               else veiculo_rotulo(v), "")
        nome = f"custo_{(v.get('placa') or veiculo_id)}.xlsx"
        return Response(
            content=buf.getvalue(), media_type=_XLSX_MEDIA,
            headers={"Content-Disposition":
                     f'attachment; filename="{nome}"'})

    @app.get("/frota/abastecimentos/{abast_id}/nfs")
    def frota_nfs_lista(abast_id: int, request: Request,
                        user: dict = auth.require_permission("frota")):
        repo = _repo()
        a = repo.get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        return templates.TemplateResponse(
            request=request, name="frota_nfs.html",
            context=ctx(request, a=a, data_br=_br(a["data"]),
                        veiculo_rotulo=veiculo_rotulo(a),
                        combustivel_label=label_combustivel(a["combustivel"]),
                        revisao_rotulo=rotulo_revisao(a.get("revisao")),
                        nfs=repo.list_nfs(abast_id),
                        assinado=repo.get_abast_signed(abast_id),
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.post("/frota/abastecimentos/{abast_id}/nfs")
    async def frota_nf_add(abast_id: int, request: Request,
                           numero: str = Form(""), data_nf: str = Form(""),
                           valor: str = Form(""),
                           arquivo: UploadFile = File(None),
                           user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user,
                        f"/frota/abastecimentos/{abast_id}/nfs")
        if red:
            return red
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione o arquivo da nota fiscal.")
            return RedirectResponse(
                f"/frota/abastecimentos/{abast_id}/nfs", status_code=303)
        data = await arquivo.read()
        try:
            val = float(str(valor).replace(",", ".")) if str(valor).strip() \
                else None
            _repo().add_nf(abast_id, numero, _iso(data_nf), val,
                           Path(arquivo.filename).name, data)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(
                f"/frota/abastecimentos/{abast_id}/nfs", status_code=303)
        _audit(request, "frota-nf-anexar", user["username"], numero or "",
               f"abast={abast_id}")
        flash(request, msg="Nota fiscal anexada.")
        return RedirectResponse(
            f"/frota/abastecimentos/{abast_id}/nfs", status_code=303)

    @app.get("/frota/nfs/{nf_id}/download")
    def frota_nf_download(nf_id: int):
        nf = _repo().get_nf(nf_id)
        if not nf:
            return Response(status_code=404)
        return Response(content=bytes(nf["dados"]),
                        media_type=_mime(nf["tipo_arquivo"]),
                        headers={"Content-Disposition":
                                 f'attachment; filename="{nf["filename"]}"'})

    @app.post("/frota/nfs/{nf_id}/excluir")
    def frota_nf_excluir(nf_id: int, request: Request,
                         user: dict = auth.require_permission("frota")):
        nf = _repo().get_nf(nf_id)
        if not nf:
            flash(request, erro="Nota fiscal não encontrada.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        url = f"/frota/abastecimentos/{nf['abastecimento_id']}/nfs"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().delete_nf(nf_id)
        flash(request, msg="Nota fiscal excluída.")
        return RedirectResponse(url, status_code=303)

    @app.post("/frota/abastecimentos/{abast_id}/valores")
    def frota_abast_valores(abast_id: int, request: Request,
                            litros: str = Form(""), valor: str = Form(""),
                            user: dict = auth.require_permission("frota")):
        """2.29.7: completa litros/valor DEPOIS (ex.: com a NF em mãos)."""
        a = _repo().get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        url = f"/frota/abastecimentos/{abast_id}/nfs"
        red = _bloqueio(request, user, url)
        if red:
            return red

        def _num(txt):
            return float(str(txt).replace(",", ".")) if str(txt).strip() \
                else None
        try:
            l, v = _num(litros), _num(valor)
            if l is not None and l <= 0:
                raise ValueError("Litros deve ser maior que zero.")
            if v is not None and v < 0:
                raise ValueError("Valor não pode ser negativo.")
            _repo().update_abast_valores(abast_id, l, v)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(url, status_code=303)
        _audit(request, "frota-abast-valores", user["username"],
               a["serial"], f"litros={l} valor={v}")
        flash(request, msg="Valores atualizados.")
        return RedirectResponse(url, status_code=303)

    @app.get("/frota/abastecimentos/{abast_id}/pdf")
    def frota_abast_pdf(abast_id: int):
        a = _repo().get_abastecimento(abast_id)
        if not a or not a.get("pdf_path"):
            return Response(status_code=404)
        p = Path(a["pdf_path"])
        if not p.exists():
            return Response(status_code=404)
        # 2.38.1: nunca servir PDF de cache — edição regenera o mesmo arquivo
        return FileResponse(p, media_type="application/pdf",
                            headers={"Cache-Control": "no-store, must-revalidate"})

    @app.get("/frota/abastecimentos/{abast_id}/pdf/download")
    def frota_abast_pdf_download(abast_id: int):
        a = _repo().get_abastecimento(abast_id)
        if not a or not a.get("pdf_path"):
            return Response(status_code=404)
        p = Path(a["pdf_path"])
        if not p.exists():
            return Response(status_code=404)
        return FileResponse(p, media_type="application/pdf",
                            filename=f"{a['serial']}.pdf",
                            headers={"Cache-Control": "no-store, must-revalidate"})


# ================= ficha do veiculo =================

    @app.get("/frota/{veiculo_id}")
    def frota_ficha(veiculo_id: int, request: Request,
                    user: dict = auth.require_permission("frota")):
        repo = _repo()
        v = repo.get_veiculo(veiculo_id)
        if not v:
            flash(request, erro="Veículo não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        v["rotulo"] = veiculo_rotulo(v)
        _status_veiculo(repo, v)  # 2.33.2: situação atual também na ficha
        docs = repo.list_docs(veiculo_id)
        laudos = repo.list_laudos(veiculo_id)
        movs = repo.list_movimentacoes(veiculo_id)
        abasts = repo.list_abast_por_veiculo(veiculo_id)
        nf_map = repo.tem_nf([a["id"] for a in abasts])
        for a in abasts:
            nf_num = nf_map.get(a["id"])
            a["nf_numero"] = nf_num
            a["tem_nf"] = bool(nf_num)
            a["bloqueada"] = (a.get("status") == "bloqueada")
        checklists = repo.list_checklists(veiculo_id)
        manutencoes = repo.list_manutencoes(veiculo_id)
        resumo = repo.resumo_custo_veiculo(veiculo_id)
        serie = repo.custo_serie_veiculo(veiculo_id)
        return templates.TemplateResponse(
            request=request, name="frota_ficha.html",
            context=ctx(request, v=v, rotulo=v["rotulo"],
                        tipo_label=label_tipo(v["tipo"]),
                        sub_label=label_subtipo(v["subtipo"]),
                        docs=docs, laudos=laudos, movs=movs, abasts=abasts,
                        checklists=checklists, manutencoes=manutencoes,
                        resumo=resumo,
                        grafico_svg=_svg_grafico_custo(serie),
                        tags_doc=TAGS_DOC,
                        funcionarios=_funcionarios_nomes(),
                        combustiveis=TIPOS_COMBUSTIVEL,
                        tipos_laudo=TIPOS_LAUDO,
                        hoje=date.today().strftime("%d/%m/%Y"),
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    # ---- documentos (pasta virtual) ----

    @app.post("/frota/{veiculo_id}/docs")
    async def frota_doc_upload(veiculo_id: int, request: Request,
                               tag: str = Form(""),
                               arquivo: UploadFile = File(None),
                               user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione um arquivo.")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        data = await arquivo.read()
        try:
            _repo().add_doc(veiculo_id, Path(arquivo.filename).name, data,
                            tag=tag)
        except ValueError as e:
            flash(request, erro=str(e))
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        flash(request, msg=f"Documento '{Path(arquivo.filename).name}'"
                           " anexado.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.get("/frota/{veiculo_id}/docs/{doc_id}/download")
    def frota_doc_download(veiculo_id: int, doc_id: int):
        doc = _repo().get_doc(doc_id)
        if not doc or doc["veiculo_id"] != veiculo_id:
            return Response(status_code=404)
        return Response(content=bytes(doc["dados"]),
                        media_type=_mime(doc["tipo"]),
                        headers={"Content-Disposition":
                                 f'attachment; filename="{doc["filename"]}"'})

    @app.post("/frota/{veiculo_id}/docs/{doc_id}/excluir")
    def frota_doc_excluir(veiculo_id: int, doc_id: int, request: Request,
                          user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        _repo().delete_doc(doc_id)
        flash(request, msg="Documento excluído.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    # ---- laudos ----

    @app.post("/frota/{veiculo_id}/laudos")
    async def frota_laudo_add(veiculo_id: int, request: Request,
                              tipo: str = Form(""),
                              descricao: str = Form(""),
                              emissao: str = Form(""),
                              validade: str = Form(""),
                              arquivo: UploadFile = File(None),
                              user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        data_emi = _iso(emissao)
        data_val = _iso(validade)
        if data_emi is None or data_val is None:
            flash(request, erro="Datas em formato inválido (dd/mm/aaaa).")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione o arquivo do laudo.")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        data = await arquivo.read()
        try:
            _repo().add_laudo(veiculo_id, tipo,
                              Path(arquivo.filename).name, data,
                              data_emi, data_val, descricao)
        except ValueError as e:
            flash(request, erro=str(e))
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        flash(request, msg="Laudo registrado.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.get("/frota/{veiculo_id}/laudos/{laudo_id}/download")
    def frota_laudo_download(veiculo_id: int, laudo_id: int):
        laudo = _repo().get_laudo(laudo_id)
        if not laudo or laudo["veiculo_id"] != veiculo_id:
            return Response(status_code=404)
        return Response(content=bytes(laudo["dados"]),
                        media_type=_mime(laudo["tipo_arquivo"]),
                        headers={"Content-Disposition":
                                 f'attachment; filename="{laudo["filename"]}"'})

    @app.post("/frota/{veiculo_id}/laudos/{laudo_id}/excluir")
    def frota_laudo_excluir(veiculo_id: int, laudo_id: int,
                            request: Request,
                            user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        _repo().delete_laudo(laudo_id)
        flash(request, msg="Laudo excluído.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    # ---- movimentacoes ----

    @app.post("/frota/{veiculo_id}/movimentacoes")
    def frota_mov_saida(veiculo_id: int, request: Request,
                        data_saida: str = Form(""), hora: str = Form(""),
                        km_inicial: str = Form(""), destino: str = Form(""),
                        motivo: str = Form(""), obs: str = Form(""),
                        motorista: str = Form(""),
                        autorizado_por: str = Form(""),
                        user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        data_iso = _iso(data_saida)
        if data_iso is None:
            flash(request, erro="Data de saída inválida (dd/mm/aaaa).")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        # 2.29.7: não permite segunda saída com uma ainda aberta
        abertas = [m for m in _repo().list_movimentacoes(veiculo_id)
                   if m.get("aberta")]
        if abertas:
            m0 = abertas[0]
            flash(request, erro="Já existe uma saída aberta neste veículo "
                  f"({_br(m0['data_saida'])} {m0['hora_saida']} — "
                  f"{m0['motorista']}). Registre a entrada antes de uma nova"
                  " saída.")
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        try:
            km = int(km_inicial) if km_inicial.strip() else None
            _repo().add_mov_saida(veiculo_id, data_iso, hora.strip(), km,
                                  destino, motivo, obs, motorista,
                                  autorizado_por)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        flash(request, msg="Saída registrada.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.post("/frota/movimentacoes/{mov_id}/entrada")
    def frota_mov_entrada(mov_id: int, request: Request,
                          data_entrada: str = Form(""),
                          hora_entrada: str = Form(""),
                          km_final: str = Form(""),
                          user: dict = auth.require_permission("frota")):
        mov = _repo().get_mov(mov_id)
        if not mov:
            flash(request, erro="Movimentação não encontrada.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{mov['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        data_iso = _iso(data_entrada)
        if data_iso is None:
            flash(request, erro="Data de entrada inválida (dd/mm/aaaa).")
            return RedirectResponse(url, status_code=303)
        try:
            km = int(km_final) if km_final.strip() else None
            _repo().registrar_entrada(mov_id, data_iso,
                                      hora_entrada.strip(), km)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(url, status_code=303)
        flash(request, msg="Entrada registrada.")
        return RedirectResponse(url, status_code=303)

    @app.post("/frota/movimentacoes/{mov_id}/excluir")
    def frota_mov_excluir(mov_id: int, request: Request,
                          user: dict = auth.require_permission("frota")):
        mov = _repo().get_mov(mov_id)
        if not mov:
            flash(request, erro="Movimentação não encontrada.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{mov['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().delete_mov(mov_id)
        flash(request, msg="Movimentação excluída.")
        return RedirectResponse(url, status_code=303)

    # ---- manutenções preventivas por KM ----

    @app.post("/frota/{veiculo_id}/manutencoes")
    def frota_manut_add(veiculo_id: int, request: Request,
                        descricao: str = Form(""), intervalo_km: str = Form(""),
                        km_ultima: str = Form(""), data_ultima: str = Form(""),
                        obs: str = Form(""),
                        user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        data_iso = _iso(data_ultima) if data_ultima.strip() else ""
        try:
            _repo().add_manutencao(
                veiculo_id, descricao, int(intervalo_km) if intervalo_km.strip()
                else 0, int(km_ultima) if km_ultima.strip() else 0,
                data_iso or "", obs)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)
        _audit(request, "frota-manutencao-criar", user["username"],
               str(veiculo_id), descricao)
        flash(request, msg="Manutenção cadastrada.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.post("/frota/manutencoes/{manut_id}/concluir")
    def frota_manut_concluir(manut_id: int, request: Request,
                             km_feito: str = Form(""), data_feito: str = Form(""),
                             user: dict = auth.require_permission("frota")):
        m = _repo().get_manutencao(manut_id)
        if not m:
            flash(request, erro="Manutenção não encontrada.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{m['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        try:
            data_iso = _iso(data_feito) or date.today().isoformat()
            _repo().concluir_manutencao(manut_id, int(km_feito), data_iso)
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(url, status_code=303)
        flash(request, msg="Manutenção registrada como realizada.")
        return RedirectResponse(url, status_code=303)

    @app.post("/frota/manutencoes/{manut_id}/excluir")
    def frota_manut_excluir(manut_id: int, request: Request,
                            user: dict = auth.require_permission("frota")):
        m = _repo().get_manutencao(manut_id)
        if not m:
            flash(request, erro="Manutenção não encontrada.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{m['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().delete_manutencao(manut_id)
        flash(request, msg="Manutenção excluída.")
        return RedirectResponse(url, status_code=303)

    # ---- checklist semanal (2.29.5) ----

    @app.get("/frota/{veiculo_id}/checklist/branco")
    def frota_checklist_branco(veiculo_id: int, request: Request,
                               user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        repo = _repo()
        v = repo.get_veiculo(veiculo_id)
        if not v:
            flash(request, erro="Veículo não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        hoje = date.today().isoformat()
        movs = repo.list_movimentacoes(veiculo_id)
        kms = [m["km_final"] if m.get("km_final") is not None else m["km_inicial"]
               for m in movs]
        km_atual = max([k for k in kms if k is not None], default=None)
        chk_id, serial = repo.add_checklist(
            veiculo_id, hoje, hoje, km_atual,
            "(a preencher)", "(a preencher)", "S", {}, {},
            criado_por=user["username"])
        try:
            from src.core.pdf_checklist import gerar_pdf_checklist
            from src.core.config import load_company_config
            from src.utils.error_log import log_error
            pdf = gerar_pdf_checklist({
                "serial": serial, "veiculo": v,
                "data_inicial_br": _br(hoje), "data_final_br": _br(hoje),
                "km_rodado": km_atual, "placa": v.get("placa"),
                "motorista": "(a preencher)", "lider": "(a preencher)",
                "pode_operar": "", "itens": {}, "observacoes": {},
                "config": load_company_config()})
            repo.set_checklist_pdf(chk_id, str(pdf))
        except Exception as e:
            log_error("portal-frota-checklist-pdf", e)
        users.audit("frota-checklist-branco", user["username"], serial)
        flash(request, msg=f"Checklist {serial} gerado para preenchimento manual.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.get("/frota/{veiculo_id}/checklist/novo")
    def frota_checklist_novo(veiculo_id: int, request: Request,
                             user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        v = _repo().get_veiculo(veiculo_id)
        if not v:
            flash(request, erro="Veículo não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        return templates.TemplateResponse(
            request=request, name="frota_checklist_form.html",
            context=ctx(request, v=v, rotulo=veiculo_rotulo(v),
                        grupos=CHECKLIST_GRUPOS, dias=CHECKLIST_DIAS,
                        hoje=date.today().strftime("%d/%m/%Y"),
                        pode_escrever=auth.pode_escrever(user["papel"],
                                                         "frota")))

    @app.post("/frota/{veiculo_id}/checklist")
    async def frota_checklist_criar(veiculo_id: int, request: Request,
                                    data_inicial: str = Form(""),
                                    data_final: str = Form(""),
                                    km_rodado: str = Form(""),
                                    motorista: str = Form(""),
                                    lider: str = Form(""),
                                    pode_operar: str = Form(""),
                                    user: dict = auth.require_permission("frota")):
        red = _bloqueio(request, user, f"/frota/{veiculo_id}")
        if red:
            return red
        url_novo = f"/frota/{veiculo_id}/checklist/novo"
        form = await request.form()
        itens: dict = {}
        observacoes: dict = {}
        for k, val in form.items():
            v_str = str(val)
            if k.startswith("item_") and v_str in ("S", "N"):
                num, _, dia = k[5:].rpartition("_")
                if num and dia:
                    itens.setdefault(num, {})[dia] = v_str
            elif k.startswith("obs_") and v_str.strip():
                observacoes[k[4:]] = v_str.strip()
        data_ini_iso = _iso(data_inicial)
        data_fim_iso = _iso(data_final)
        if data_ini_iso is None or data_fim_iso is None:
            flash(request, erro="Datas inválidas (dd/mm/aaaa).")
            return RedirectResponse(url_novo, status_code=303)
        try:
            km = int(km_rodado) if km_rodado.strip() else None
            chk_id, serial = _repo().add_checklist(
                veiculo_id, data_ini_iso, data_fim_iso, km,
                motorista, lider, pode_operar, itens, observacoes,
                user["username"])
        except (ValueError, TypeError) as e:
            flash(request, erro=str(e))
            return RedirectResponse(url_novo, status_code=303)
        try:
            from src.core.pdf_checklist import gerar_pdf_checklist
            from src.core.config import load_company_config
            v = _repo().get_veiculo(veiculo_id)
            pdf = gerar_pdf_checklist({
                "serial": serial,
                "veiculo": v,
                "data_inicial_br": _br(data_ini_iso),
                "data_final_br": _br(data_fim_iso),
                "km_rodado": km,
                "placa": v.get("placa"),
                "motorista": motorista,
                "lider": lider,
                "pode_operar": pode_operar,
                "itens": itens,
                "observacoes": observacoes,
                "config": load_company_config(),
            })
            _repo().set_checklist_pdf(chk_id, str(pdf))
        except Exception as e:
            from src.utils.error_log import log_error
            log_error("portal-frota-checklist-pdf", e)
        _audit(request, "frota-checklist", user["username"],
               serial, f"veiculo={veiculo_id}")
        flash(request, msg=f"Checklist {serial} registrado.")
        return RedirectResponse(f"/frota/{veiculo_id}", status_code=303)

    @app.get("/frota/checklists/{chk_id}/pdf")
    def frota_checklist_pdf(chk_id: int):
        c = _repo().get_checklist(chk_id)
        if not c or not c.get("pdf_path"):
            return Response(status_code=404)
        p = Path(c["pdf_path"])
        if not p.exists():
            return Response(status_code=404)
        return FileResponse(p, media_type="application/pdf")

    @app.get("/frota/checklists/{chk_id}/pdf/download")
    def frota_checklist_pdf_download(chk_id: int):
        c = _repo().get_checklist(chk_id)
        if not c or not c.get("pdf_path"):
            return Response(status_code=404)
        p = Path(c["pdf_path"])
        if not p.exists():
            return Response(status_code=404)
        return FileResponse(p, media_type="application/pdf",
                            filename=f"{c['serial']}.pdf")

    @app.post("/frota/checklists/{chk_id}/assinado")
    async def frota_checklist_assinado(chk_id: int, request: Request,
                                       arquivo: UploadFile = File(None),
                                       user: dict = auth.require_permission(
                                           "frota")):
        """2.29.7: fluxo papel — imprime, preenche, assina e ANEXA de volta."""
        c = _repo().get_checklist(chk_id)
        if not c:
            flash(request, erro="Checklist não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{c['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione o arquivo assinado (PDF, JPG ou"
                                " PNG).")
            return RedirectResponse(url, status_code=303)
        ext = Path(arquivo.filename).suffix.lower().lstrip(".")
        tipo = {"pdf": "pdf", "jpg": "jpg", "jpeg": "jpg",
                "png": "png"}.get(ext)
        data = await arquivo.read()
        if tipo is None:
            flash(request, erro="Formato não suportado (use PDF, JPG ou"
                                " PNG).")
            return RedirectResponse(url, status_code=303)
        if len(data) > 50 * 1024 * 1024:
            flash(request, erro="Arquivo maior que 50MB.")
            return RedirectResponse(url, status_code=303)
        _repo().attach_checklist_signed(
            chk_id, data, tipo, Path(arquivo.filename).name)
        _audit(request, "frota-checklist-assinado", user["username"],
               c["serial"], Path(arquivo.filename).name)
        flash(request, msg="Checklist assinado anexado.")
        return RedirectResponse(url, status_code=303)

    @app.get("/frota/checklists/{chk_id}/assinado/download")
    def frota_checklist_assinado_download(chk_id: int):
        s = _repo().get_checklist_signed(chk_id)
        if not s:
            return Response(status_code=404)
        media = {"pdf": "application/pdf", "jpg": "image/jpeg",
                 "png": "image/png"}.get(s["assinado_tipo"],
                                         "application/octet-stream")
        return Response(content=bytes(s["assinado_dados"]), media_type=media,
                        headers={"Content-Disposition":
                                 f'attachment; '
                                 f'filename="{s["assinado_filename"]}"'})

    @app.post("/frota/checklists/{chk_id}/assinado/excluir")
    def frota_checklist_assinado_excluir(chk_id: int, request: Request,
                                         user: dict = auth.require_permission(
                                             "frota")):
        c = _repo().get_checklist(chk_id)
        if not c:
            flash(request, erro="Checklist não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{c['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().remove_checklist_signed(chk_id)
        flash(request, msg="Anexo assinado removido.")
        return RedirectResponse(url, status_code=303)

    @app.post("/frota/abastecimentos/{abast_id}/assinado")
    async def frota_abast_assinado(abast_id: int, request: Request,
                                   arquivo: UploadFile = File(None),
                                   user: dict = auth.require_permission(
                                       "frota")):
        """2.31: abastecimento assinado anexado de volta (foto ou PDF)."""
        a = _repo().get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        url = f"/frota/abastecimentos/{abast_id}/nfs"
        red = _bloqueio(request, user, url)
        if red:
            return red
        if arquivo is None or not arquivo.filename:
            flash(request, erro="Selecione o arquivo assinado (PDF, JPG ou"
                                " PNG).")
            return RedirectResponse(url, status_code=303)
        ext = Path(arquivo.filename).suffix.lower().lstrip(".")
        tipo = {"pdf": "pdf", "jpg": "jpg", "jpeg": "jpg",
                "png": "png"}.get(ext)
        data = await arquivo.read()
        if tipo is None:
            flash(request, erro="Formato não suportado (use PDF, JPG ou"
                                " PNG).")
            return RedirectResponse(url, status_code=303)
        if len(data) > 50 * 1024 * 1024:
            flash(request, erro="Arquivo maior que 50MB.")
            return RedirectResponse(url, status_code=303)
        _repo().attach_abast_signed(
            abast_id, data, tipo, Path(arquivo.filename).name)
        _audit(request, "frota-abast-assinado", user["username"],
               a["serial"], Path(arquivo.filename).name)
        flash(request, msg="Abastecimento assinado anexado.")
        return RedirectResponse(url, status_code=303)

    @app.get("/frota/abastecimentos/{abast_id}/assinado/download")
    def frota_abast_assinado_download(abast_id: int):
        s = _repo().get_abast_signed(abast_id)
        if not s:
            return Response(status_code=404)
        media = {"pdf": "application/pdf", "jpg": "image/jpeg",
                 "png": "image/png"}.get(s["assinado_tipo"],
                                         "application/octet-stream")
        return Response(content=bytes(s["assinado_dados"]), media_type=media,
                        headers={"Content-Disposition":
                                 f'attachment; '
                                 f'filename="{s["assinado_filename"]}"'})

    @app.post("/frota/abastecimentos/{abast_id}/assinado/excluir")
    def frota_abast_assinado_excluir(abast_id: int, request: Request,
                                     user: dict = auth.require_permission(
                                         "frota")):
        a = _repo().get_abastecimento(abast_id)
        if not a:
            flash(request, erro="Abastecimento não encontrado.")
            return RedirectResponse("/frota/abastecimentos", status_code=303)
        url = f"/frota/abastecimentos/{abast_id}/nfs"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().remove_abast_signed(abast_id)
        flash(request, msg="Anexo assinado removido.")
        return RedirectResponse(url, status_code=303)

    @app.post("/frota/checklists/{chk_id}/excluir")
    def frota_checklist_excluir(chk_id: int, request: Request,
                                user: dict = auth.require_permission("frota")):
        c = _repo().get_checklist(chk_id)
        if not c:
            flash(request, erro="Checklist não encontrado.")
            return RedirectResponse("/frota", status_code=303)
        url = f"/frota/{c['veiculo_id']}"
        red = _bloqueio(request, user, url)
        if red:
            return red
        _repo().delete_checklist(chk_id)
        flash(request, msg="Checklist excluído.")
        return RedirectResponse(url, status_code=303)

    
