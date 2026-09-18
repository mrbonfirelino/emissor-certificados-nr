"""
Certificados PPTX da tecnica de seguranca (NR-06, NR-12, NR-18, NR-35).

Fluxo:
1. Modelos preparados ficam em templates/certificados_pptx/NR-XX.pptx
   (gerados por tools/prepare_nr_pptx.py a partir de "MODELOS NR pptx\").
2. O modelo contem tokens {{NOME}} {{CPF}} {{DIA}} {{MES}} {{ANO}} — a
   substituicao preserva a formatacao por run (o token herda o estilo do
   primeiro run afetado).
3. Em todos os slides e adicionada uma caixa pequena cinza no canto inferior
   direito com o numero do certificado (ex.: CERT-000491).
4. Conversao para PDF via PowerPoint COM (reutiliza pptx_card_service).
"""

import re
import shutil
import tempfile
from datetime import date
from pathlib import Path
from typing import Dict, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.utils.paths import get_templates_dir

MESES_PT = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

TOKEN_RE = re.compile(r"\{\{([A-Z_0-9]+)\}\}")


def get_pptx_cert_dir() -> Path:
    return Path(get_templates_dir()) / "certificados_pptx"


def get_pptx_template_path(nr_code: str) -> Optional[Path]:
    p = get_pptx_cert_dir() / f"{nr_code}.pptx"
    return p if p.exists() else None


def valores_certificado(employee, data_treinamento: date) -> Dict[str, str]:
    d = data_treinamento
    return {
        "NOME": (employee.nome or "").upper(),
        "CPF": (employee.cpf or "").strip(),
        "DIA": str(d.day),
        "MES": MESES_PT[d.month - 1],
        "ANO": str(d.year),
    }


def _substituir_no_paragrafo(para, values: Dict[str, str]) -> bool:
    """
    Substitui {{TOKENS}} no paragrafo preservando a formatacao por run:
    o valor substitui o trecho do token herdando o estilo do primeiro run
    afetado; prefixos/sufixos mantem o estilo de cada run.
    """
    runs = para.runs
    if not runs:
        return False
    texts = [r.text or "" for r in runs]
    full = "".join(texts)
    matches = list(TOKEN_RE.finditer(full))
    if not matches:
        return False

    bounds = []
    pos = 0
    for t in texts:
        bounds.append((pos, pos + len(t)))
        pos += len(t)

    for m in reversed(matches):
        start, end = m.span()
        repl = values.get(m.group(1), m.group(0))
        i0 = next(i for i, (s, e) in enumerate(bounds) if s <= start < e)
        i1 = next((i for i, (s, e) in enumerate(bounds) if s < end <= e), i0)
        pre = texts[i0][: start - bounds[i0][0]]
        if i0 == i1:
            texts[i0] = pre + repl + texts[i0][end - bounds[i0][0]:]
        else:
            texts[i0] = pre + repl
            for i in range(i0 + 1, i1):
                texts[i] = ""
            texts[i1] = texts[i1][end - bounds[i1][0]:]

    for run, novo in zip(runs, texts):
        if run.text != novo:
            run.text = novo
    return True


def _iterar_text_frames(shapes):
    """Text frames de shapes (inclusive grupos) e celulas de tabela."""
    stack = list(shapes)
    while stack:
        shape = stack.pop()
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            stack.extend(shape.shapes)
            continue
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
            continue
        if getattr(shape, "has_text_frame", False):
            yield shape.text_frame


def _add_caixa_numero(prs, slide, cert_number: str, data_hora: str = ""):
    """Caixa pequena cinza no canto inferior direito com o numero do certificado
    (e, opcionalmente, a data/hora da emissao)."""
    w = int(prs.slide_width or 0)
    h = int(prs.slide_height or 0)
    box = slide.shapes.add_textbox(
        w - Inches(2.35), h - Inches(0.42), Inches(2.1), Inches(0.3)
    )
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = Inches(0.08)
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{cert_number} • {data_hora}" if data_hora else cert_number
    f = run.font
    f.size = Pt(9)
    f.name = "Arial"
    f.color.rgb = RGBColor(0x8C, 0x8C, 0x8C)


def gerar_pdf_pptx(
    nr_code: str,
    employee,
    data_treinamento: date,
    cert_number: str,
    pdf_path: Path,
    data_hora: str = "",
) -> Path:
    """
    Preenche o modelo PPTX da NR e converte para PDF (PowerPoint COM).
    Salva em pdf_path (mesmo padrao de pasta/nome do fluxo JSON/ReportLab).
    """
    tpl = get_pptx_template_path(nr_code)
    if not tpl:
        raise ValueError(f"Modelo PPTX {nr_code} nao encontrado em templates/certificados_pptx/")

    prs = Presentation(str(tpl))
    values = valores_certificado(employee, data_treinamento)
    for slide in prs.slides:
        for tf in _iterar_text_frames(slide.shapes):
            for para in tf.paragraphs:
                _substituir_no_paragrafo(para, values)
        _add_caixa_numero(prs, slide, cert_number, data_hora)

    tmp = Path(tempfile.mkdtemp(prefix="cert_pptx_"))
    try:
        clone = tmp / "certificado.pptx"
        prs.save(str(clone))
        from src.core.pptx_card_service import _pptx_to_pdf_batch

        out_pdf = clone.with_suffix(".pdf")
        _pptx_to_pdf_batch([(clone, out_pdf)])
        pdf_path = Path(pdf_path)
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(out_pdf, pdf_path)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return pdf_path
