"""
Gerador de Lista de Presenca (v1.41.0).

- Templates reais (XLSX com a logo embutida) em templates/listas_presenca/,
  mapeados por registry.json (NR-01, NR-06, NR-12, NR-18, NR-35).
- Preenchimento via openpyxl (deterministico, sem COM): edita data, carga e
  participantes (empresa sempre ALTEC INDUSTRIAL), LIMPA as vagas nao
  utilizadas (dados de exemplo do template), remove as demais abas, define
  area de impressao explicita (sem folhas em branco), folha VERTICAL e
  substitui os tokens {SERIAL} e {PAGINACAO} do modelo.
- Muitos participantes: o mesmo modelo e preenchido em folhas separadas
  (Pag i de N) e os PDFs sao mesclados.
- Conversao para PDF: Microsoft Excel (COM), apenas ExportAsFixedFormat.
- NRs sem modelo proprio: layout padrao no estilo da NR 01 (gerado com
  openpyxl, com a logo da empresa no topo).
- Participantes: puxados das emissoes de certificados daquela NR naquela data
  (HistoryRepository.query) — veja participantes_da_emissao().
"""

import json
import re
import shutil
import tempfile
from datetime import date
from pathlib import Path
from typing import List, Optional, Tuple

from src.utils.paths import get_templates_dir, get_data_dir, get_logo_path

REGISTRY_ARQ = "registry.json"

EMPRESA_PADRAO = "ALTEC INDUSTRIAL"
VAGAS_PADRAO = 20  # participantes por folha quando o modelo nao declara vagas

XL_TYPE_PDF = 0  # Excel: xlTypePDF

_DOC_EXT_OK = {"pdf": "application/pdf", "jpg": "image/jpeg",
               "jpeg": "image/jpeg", "png": "image/png"}


def get_listas_dir() -> Path:
    return get_data_dir() / "listas_presenca"


def get_registry_dir() -> Path:
    return Path(get_templates_dir()) / "listas_presenca"


def _carregar_registry() -> dict:
    f = get_registry_dir() / REGISTRY_ARQ
    if not f.exists():
        return {}
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except Exception:
        return {}


def tipos_com_modelo() -> List[str]:
    """NRs que possuem template XLSX real (registry)."""
    return sorted(_carregar_registry().keys())


def _ext_ok(nome: str) -> Optional[str]:
    ext = Path(nome).suffix.lower().lstrip(".")
    return _DOC_EXT_OK.get(ext)


# ── Participantes puxados da emissao ───────────────────────

def participantes_da_emissao(nr_code: str, data_ref: str,
                             carga_minima: Optional[float] = None) -> Tuple[list, float]:
    """
    Lista de participantes ({nome, funcao, cpf}) dos certificados daquela NR
    com data_fim == data_ref (ISO), e a carga horaria usada nas emissoes do
    dia (a maior; fallback: carga_minima do template ou 0).
    """
    from src.core.history_repo import HistoryRepository
    from src.core.employee_repo import EmployeeRepository

    recs = HistoryRepository().query(nr_code=nr_code, data_de=data_ref, data_ate=data_ref,
                                     limit=10000)
    er = EmployeeRepository()
    funcs_cache = {}
    vistos = set()
    out = []
    cargas = []
    for rec in recs:
        try:
            cargas.append(float(rec.carga_horaria or 0))
        except (TypeError, ValueError):
            pass
        chave = (rec.funcionario_nome or "").strip().upper()
        if not chave or chave in vistos:
            continue
        vistos.add(chave)
        funcao = ""
        if rec.employee_id:
            if rec.employee_id not in funcs_cache:
                emp = er.get_by_id(rec.employee_id)
                funcs_cache[rec.employee_id] = (emp.funcao or "") if emp else ""
            funcao = funcs_cache[rec.employee_id]
        out.append({"nome": rec.funcionario_nome or "",
                    "funcao": funcao, "cpf": rec.funcionario_cpf or ""})
    carga = max(cargas) if cargas else float(carga_minima or 0)
    return out, carga


# ── Paginacao e tokens (openpyxl) ───────────────────────────

def _paginacao_openpyxl(ws, ultima_col: int, ultima_linha: int):
    """
    Folha VERTICAL (retrato), 1 pagina de largura, grade visivel, centrado e
    AREA DE IMPRESSAO explicita (sem ela, celulas formatadas vazias viram
    folhas em branco no PDF).
    """
    try:
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.properties import PageSetupProperties

        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.gridLines = True
        ws.print_options.horizontalCentered = True
        ws.print_area = f"A1:{get_column_letter(max(ultima_col, 1))}{max(ultima_linha, 1)}"
    except Exception:
        pass


# ── Conversao XLSX -> PDF (Excel COM) ──────────────────────

def _excel_app():
    """
    Cria o Excel COM ja com CoInitialize: handlers web (sync def) rodam em
    threads do anyio SEM COM inicializado — sem isso da
    'CoInitialize nao foi chamado' (-2147221008). Use com _fechar_excel().
    """
    try:
        import comtypes.client
        import comtypes
    except ImportError:
        raise RuntimeError("Biblioteca comtypes nao instalada. "
                           "Instale com: pip install comtypes")
    comtypes.CoInitialize()
    try:
        excel = comtypes.client.CreateObject("Excel.Application", dynamic=True)
    except Exception as e:
        try:
            import comtypes

            comtypes.CoUninitialize()
        except Exception:
            pass
        raise RuntimeError(
            "Microsoft Excel nao encontrado neste computador.\n"
            "As listas de presenca exigem o Excel instalado (Office).\n"
            f"Detalhe: {e}"
        ) from e
    excel.Visible = False
    excel.DisplayAlerts = False
    return excel


def _fechar_excel(excel, wb=None):
    try:
        if wb is not None:
            wb.Close(False)
    except Exception:
        pass
    try:
        excel.Quit()
    except Exception:
        pass
    try:
        import comtypes

        comtypes.CoUninitialize()
    except Exception:
        pass


def _excel_to_pdf(xlsx_path: Path, pdf_path: Path):
    excel = _excel_app()
    wb = None
    try:
        wb = excel.Workbooks.Open(str(Path(xlsx_path).resolve()))
        wb.ExportAsFixedFormat(XL_TYPE_PDF, str(Path(pdf_path).resolve()))
    finally:
        _fechar_excel(excel, wb)


def _celula(ws, ref: str):
    """
    Celula escrevivel: se ref cair dentro de um intervalo mesclado, devolve a
    celula-ancora (openpyxl recusa escrita em MergedCell).
    """
    try:
        from openpyxl.utils.cell import coordinate_to_tuple

        linha, col = coordinate_to_tuple(ref)
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= linha <= rng.max_row and rng.min_col <= col <= rng.max_col:
                return ws.cell(row=rng.min_row, column=rng.min_col)
    except Exception:
        pass
    return ws[ref]


def _substituir_tokens(ws, serial: str, pagina: int, total: int):
    """Substitui {SERIAL} e {PAGINACAO} nos textos do modelo (ex.: C4/D4)."""
    pag = f"Pag {pagina} de {total}"
    for row in ws.iter_rows():
        for cel in row:
            v = cel.value
            if not isinstance(v, str):
                continue
            if "{SERIAL}" not in v and "{PAGINACAO}" not in v:
                continue
            novo = v.replace("{SERIAL}", serial).replace("{PAGINACAO}", pag)
            try:
                _celula(ws, cel.coordinate).value = novo
            except Exception:
                pass


def _preencher_modelo(caminho_xlsx: Path, cfg: dict, data_br: str, carga: float,
                      participantes: List[dict], tmp: Path,
                      serial: str = "", pagina: int = 1, total: int = 1) -> Path:
    """
    Copia o XLSX e preenche com openpyxl (sem COM — o binding de escrita do
    Excel varia por maquina/versao e ja falhou em producao): escolhe a aba,
    remove as outras, escreve data/carga/linhas (empresa sempre ALTEC
    INDUSTRIAL), limpa as vagas nao usadas (dados de exemplo do template),
    substitui os tokens {SERIAL}/{PAGINACAO} e define a area de impressao.
    Devolve o caminho do clone salvo. A logo JA VEM do modelo (nao duplica).
    """
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter

    clone = tmp / caminho_xlsx.name
    shutil.copyfile(caminho_xlsx, clone)

    wb = openpyxl.load_workbook(str(clone))
    alvo = (cfg.get("sheet") or "").strip()
    nome_ws = next((n for n in wb.sheetnames if n.strip() == alvo), None)
    if nome_ws is None:
        raise ValueError(f"Aba '{cfg['sheet']}' nao encontrada no modelo")
    ws = wb[nome_ws]
    for n in [x for x in wb.sheetnames if x != nome_ws]:
        wb.remove(wb[n])

    # tabela de conversao col->linha maxima usada nas configs (p/ print_area)
    colunas_usadas = [1]
    ultima_linha_cfg = 1

    def _marcar_col(ref: Optional[str]):
        nonlocal ultima_linha_cfg
        if not ref:
            return
        try:
            colunas_usadas.append(column_index_from_string(re.match(r"[A-Z]+", ref).group(0)))
            linha = int(re.search(r"\d+", ref).group(0))
            ultima_linha_cfg = max(ultima_linha_cfg, linha)
        except Exception:
            pass

    dcfg = cfg.get("data") or {}
    if dcfg.get("cel"):
        _marcar_col(dcfg["cel"])
        cel = _celula(ws, dcfg["cel"])
        if dcfg.get("data_pura"):
            cel.number_format = "@"  # texto: evita virar mm/dd/yyyy
        cel.value = f"{dcfg.get('prefixo', '')}{data_br}"

    ccfg = cfg.get("carga") or {}
    if ccfg.get("cel"):
        _marcar_col(ccfg["cel"])
        # carga vem do banco (REAL) e pode chegar como 8.0 — formatos
        # {:02d} exigem int; 8.5 continua decimal (fallback formato simples)
        carga_fmt = int(carga) if float(carga) == int(carga) else carga
        try:
            carga_txt = ccfg.get("formato", "{}").format(carga_fmt)
        except (ValueError, TypeError):
            carga_txt = f"{carga_fmt}HS"
        _celula(ws, ccfg["cel"]).value = carga_txt

    # carga_texto (regex sobre o conteudo programatico) NAO e mais aplicada:
    # o texto do template (ex.: "CONTEUDO PROGRAMATICO - 03 HORAS") fica
    # intacto, igual aos certificados PPTX.

    lcfg = cfg.get("linhas") or {}
    inicio = int(lcfg.get("inicio", 1))
    vagas = int(lcfg.get("vagas", 0))
    for i, p in enumerate(participantes):
        linha = inicio + i
        for chave in ("num", "nome", "funcao", "empresa"):
            ref = lcfg.get(chave)
            if not ref:
                continue
            _marcar_col(f"{ref}{linha}")
            if chave == "num":
                _celula(ws, f"{ref}{linha}").value = i + 1
            elif chave == "empresa":
                _celula(ws, f"{ref}{linha}").value = EMPRESA_PADRAO
            else:
                _celula(ws, f"{ref}{linha}").value = p.get(chave, "")
    # limpa vagas nao usadas (dados de exemplo que ja vieram no template)
    for i in range(len(participantes), vagas):
        linha = inicio + i
        for chave in ("num", "nome", "funcao", "empresa"):
            ref = lcfg.get(chave)
            if ref and ws[f"{ref}{linha}"].value not in (None, ""):
                _celula(ws, f"{ref}{linha}").value = None

    ultima_linha = max(inicio + max(vagas, len(participantes), 1),
                       ultima_linha_cfg)
    _substituir_tokens(ws, serial, pagina, total)
    _paginacao_openpyxl(ws, max(colunas_usadas), ultima_linha)

    wb.save(str(clone))
    return clone


# ── Layout padrao (NRs sem modelo proprio) ─────────────────

def _gerar_padrao_xlsx(tmp: Path, nr_label: str, data_br: str, carga: float,
                       participantes: List[dict], assunto: str,
                       serial: str = "", pagina: int = 1, total: int = 1) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    try:
        from src.core.config import load_company_config

        cfg = load_company_config()
    except Exception:
        cfg = None
    local = (getattr(cfg, "empresa_nome", "") or EMPRESA_PADRAO).upper()
    instrutor = getattr(cfg, "instrutor_nome", "") or ""
    registro = getattr(cfg, "instrutor_registro_mte", "") or ""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.sheet_view.showGridLines = False

    fino = Side(style="thin")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    centro = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = "LISTA DE PRESENÇA"
    c.font = Font(bold=True, size=18)
    c.alignment = centro

    # logo da empresa (se existir)
    logo = Path(get_logo_path()) if get_logo_path() else None
    if logo and logo.exists():
        try:
            from openpyxl.drawing.image import Image as XlImage

            img = XlImage(str(logo))
            escala = 110.0 / max(img.width, 1)
            img.width = int(img.width * escala)
            img.height = int(img.height * escala)
            img.anchor = "A2"
            ws.add_image(img)
            ws.row_dimensions[2].height = 46
        except Exception:
            pass

    ws["A4"] = "ASSUNTO:"
    ws["A4"].font = Font(bold=True)
    ws.merge_cells("B4:I4")
    ws["B4"] = assunto or nr_label
    ws["B4"].font = Font(bold=True, size=11)

    ws["A5"] = "INSTRUTOR:"
    ws["A5"].font = Font(bold=True)
    ws.merge_cells("B5:F5")
    ws["B5"] = f"{instrutor}    REGISTRO MTE {registro}".strip()
    ws["G5"] = "DATA:"
    ws["G5"].font = Font(bold=True)
    ws["H5"] = data_br
    ws["I5"] = f"Carga: {int(carga)}h" if carga == int(carga) else f"Carga: {carga}h"

    ws["A6"] = "Local de realização:"
    ws["A6"].font = Font(bold=True)
    ws.merge_cells("B6:I6")
    ws["B6"] = local

    ws["A7"] = "LISTA:"
    ws["A7"].font = Font(bold=True)
    ws.merge_cells("B7:D7")
    ws["B7"] = serial
    ws["G7"] = "PÁGINA:"
    ws["G7"].font = Font(bold=True)
    ws.merge_cells("H7:I7")
    ws["H7"] = f"Pag {pagina} de {total}"

    header = ["Nº", "EMPRESA", "NOME DO EMPREGADO", "", "FUNÇÃO", "", "", "ASSINATURA", ""]
    larguras = [5, 24, 42, 4, 24, 4, 3, 32, 3]
    for i, (titulo, larg) in enumerate(zip(header, larguras), start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg

    linha_h = 9
    ws.merge_cells(f"A{linha_h}:A{linha_h + 1}")
    ws.merge_cells(f"B{linha_h}:B{linha_h + 1}")
    ws.merge_cells(f"C{linha_h}:D{linha_h + 1}")
    ws.merge_cells(f"E{linha_h}:G{linha_h + 1}")
    ws.merge_cells(f"H{linha_h}:I{linha_h + 1}")
    for col, titulo in (("A", "Nº"), ("B", "EMPRESA"), ("C", "NOME DO EMPREGADO"),
                        ("E", "FUNÇÃO"), ("H", "ASSINATURA")):
        cel = ws[f"{col}{linha_h}"]
        cel.value = titulo
        cel.font = Font(bold=True, size=10)
        cel.alignment = centro
    for r in range(linha_h, linha_h + 2):
        for i in range(1, 10):
            ws.cell(row=r, column=i).border = borda

    linha = linha_h + 2
    for i, p in enumerate(participantes, start=1):
        ws.merge_cells(f"C{linha}:D{linha}")
        ws.merge_cells(f"E{linha}:G{linha}")
        ws.merge_cells(f"H{linha}:I{linha}")
        ws.cell(row=linha, column=1, value=i).alignment = centro
        ws.cell(row=linha, column=2, value=EMPRESA_PADRAO)
        ws.cell(row=linha, column=3, value=p.get("nome", ""))
        ws.cell(row=linha, column=5, value=p.get("funcao", ""))
        for col_i in range(1, 10):
            ws.cell(row=linha, column=col_i).border = borda
            ws.cell(row=linha, column=col_i).alignment = Alignment(vertical="center",
                                                                   wrap_text=False)
        ws.row_dimensions[linha].height = 20
        linha += 1

    # area de impressao explicita: sem folhas em branco no PDF
    _paginacao_openpyxl(ws, 9, linha)

    clone = tmp / "lista_padrao.xlsx"
    wb.save(str(clone))
    return clone


def gerar_pdf_lista(nr_code: str, nr_label: str, data_ref: str, carga: float,
                    participantes: List[dict], assunto: str,
                    saida_pdf: Path, serial: str = "") -> Path:
    """
    Gera o PDF da lista de presenca em saida_pdf.
    Usa o template real quando a NR esta no registry; senao, layout padrao.
    Muitos participantes: um clone por folha (vagas do registry; padrao
    VAGAS_PADRAO) com Pag i de N, mesclados em um unico PDF.
    O preenchimento e openpyxl; exige Microsoft Excel (COM) apenas para a
    conversao final em PDF.
    """
    if not participantes:
        raise ValueError("Nenhum participante para a lista")
    d = date.fromisoformat(data_ref)
    data_br = f"{d.day:02d}/{d.month:02d}/{d.year}"

    tmp = Path(tempfile.mkdtemp(prefix="lista_presenca_"))
    try:
        cfg = _carregar_registry().get(nr_code)
        por_folha = VAGAS_PADRAO
        if cfg:
            origem = get_registry_dir() / cfg["arquivo"]
            if not origem.exists():
                raise ValueError(f"Modelo de lista nao encontrado: {cfg['arquivo']}")
            try:
                por_folha = int((cfg.get("linhas") or {}).get("vagas", VAGAS_PADRAO))
            except (TypeError, ValueError):
                por_folha = VAGAS_PADRAO
            por_folha = max(por_folha, 1)

        folhas = []
        total = max(1, (len(participantes) + por_folha - 1) // por_folha)
        for pag, i in enumerate(range(0, len(participantes), por_folha), start=1):
            parte = participantes[i:i + por_folha]
            if cfg:
                clone = _preencher_modelo(origem, cfg, data_br, carga, parte, tmp,
                                          serial=serial, pagina=pag, total=total)
            else:
                clone = _gerar_padrao_xlsx(tmp, nr_label, data_br, carga, parte,
                                           assunto, serial=serial,
                                           pagina=pag, total=total)
            pdf_folha = tmp / f"folha_{pag}.pdf"
            _excel_to_pdf(clone, pdf_folha)
            if not pdf_folha.exists():
                raise RuntimeError("Excel nao gerou o PDF da lista")
            folhas.append(pdf_folha)

        saida_pdf = Path(saida_pdf)
        saida_pdf.parent.mkdir(parents=True, exist_ok=True)
        if len(folhas) == 1:
            shutil.copyfile(folhas[0], saida_pdf)
        else:
            import fitz

            doc = fitz.open()
            try:
                for f in folhas:
                    src = fitz.open(str(f))
                    try:
                        doc.insert_pdf(src)
                    finally:
                        src.close()
                doc.save(str(saida_pdf))
            finally:
                doc.close()
        return saida_pdf
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
