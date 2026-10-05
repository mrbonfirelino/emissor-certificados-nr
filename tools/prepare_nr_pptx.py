r"""
Prepara os modelos PPTX de certificados da tecnica de seguranca.

Pega os arquivos de "MODELOS NR pptx\" (originais, intocados), substitui os
dados de exemplo por tokens e salva em templates/certificados_pptx/:

    NR 06 - CERTIFICADO.pptx  -> templates/certificados_pptx/NR-06.pptx
    NR 12 - CERTIFICADO.pptx  -> templates/certificados_pptx/NR-12.pptx
    NR 18  - CERTIFICADO.pptx -> templates/certificados_pptx/NR-18.pptx
    NR 35 - CERTIFICADO.pptx  -> templates/certificados_pptx/NR-35.pptx

Tokens criados (o design/texto fixo NAO e alterado):
    {{NOME}}  nome do funcionario
    {{CPF}}   CPF do funcionario
    {{DIA}} {{MES}} {{ANO}}  todas as datas (corpo e assinatura) por extenso

Uso:  python tools\prepare_nr_pptx.py
Rodar de novo e seguro: sempre parte dos originais.
"""

import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "MODELOS NR pptx"
DESTINO = RAIZ / "templates" / "certificados_pptx"

# dados de exemplo presentes nos modelos -> tokens
LITERAIS = {
    "JONATAS PEREIRA CYPRIANO": "{{NOME}}",
    "FELIPE DOS REIS MATTOS": "{{NOME}}",
    "14464636721": "{{CPF}}",
    "11791119743": "{{CPF}}",
}

DATA_RE = r"\b\d{1,2}\s+de\s+[a-zà-ú]+\s+de\s+\d{4}\b"
DATA_TOKEN = "{{DIA}} de {{MES}} de {{ANO}}"

_COMBINED_RE = re.compile(
    "|".join([re.escape(l) for l in LITERAIS] + [DATA_RE]),
    re.IGNORECASE,
)


def _repl(m: "re.Match") -> str:
    s = m.group(0)
    if re.fullmatch(r"\d{11}", s):
        return "{{CPF}}"
    if re.search(r"\d{1,2}\s+de\s+", s, re.IGNORECASE):
        return DATA_TOKEN
    return LITERAIS.get(s, s)


def _replace_spans_in_paragraph(para) -> int:
    """
    Substitui ocorrencias no texto do paragrafo preservando a formatacao
    por run: o trecho substituido herda o estilo do primeiro run afetado;
    prefixos/sufixos mantem o estilo de cada run.
    Retorna quantidade de substituicoes.
    """
    runs = para.runs
    if not runs:
        return 0
    texts = [r.text or "" for r in runs]
    full = "".join(texts)
    matches = list(_COMBINED_RE.finditer(full))
    if not matches:
        return 0

    # limites (start, end) de cada run no texto original
    bounds = []
    pos = 0
    for t in texts:
        bounds.append((pos, pos + len(t)))
        pos += len(t)

    # processa de tras para frente para nao invalidar offsets anteriores
    for m in reversed(matches):
        start, end = m.span()
        repl = _repl(m)
        # run do inicio
        i0 = next(i for i, (s, e) in enumerate(bounds) if s <= start < e)
        # run do fim (end pode cair no limite exato de um run)
        i1 = next((i for i, (s, e) in enumerate(bounds) if s < end <= e), i0)
        pre = texts[i0][: start - bounds[i0][0]]
        if i0 == i1:
            suf = texts[i0][end - bounds[i0][0]:]
            texts[i0] = pre + repl + suf
        else:
            texts[i0] = pre + repl
            for i in range(i0 + 1, i1):
                texts[i] = ""
            texts[i1] = texts[i1][end - bounds[i1][0]:]

    for run, novo in zip(runs, texts):
        if run.text != novo:
            run.text = novo
    return len(matches)


def _iterar_text_frames(prs):
    """Text frames de shapes (inclusive grupos) e celulas de tabela."""
    for slide in prs.slides:
        stack = list(slide.shapes)
        while stack:
            shape = stack.pop()
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack.extend(shape.shapes)
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame
                continue
            if getattr(shape, "has_text_frame", False):
                yield shape.text_frame


def tokenizar(pptx_in: Path, pptx_out: Path) -> dict:
    prs = Presentation(str(pptx_in))
    total = 0
    for tf in _iterar_text_frames(prs):
        for para in tf.paragraphs:
            total += _replace_spans_in_paragraph(para)
    pptx_out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_out))
    return {"substituicoes": total}


def texto_completo(pptx: Path) -> str:
    prs = Presentation(str(pptx))
    partes = []
    for tf in _iterar_text_frames(prs):
        partes.append(tf.text)
    return "\n".join(partes)


def main() -> int:
    if not ORIGEM.exists():
        print(f"[ERRO] Pasta de origem nao encontrada: {ORIGEM}")
        return 1

    arquivos = sorted(ORIGEM.glob("*.pptx"))
    if not arquivos:
        print(f"[ERRO] Nenhum .pptx em {ORIGEM}")
        return 1

    ok = True
    for src in arquivos:
        m = re.search(r"NR\s*(\d+)", src.stem, re.IGNORECASE)
        if not m:
            print(f"[AVISO] Ignorado (sem numero de NR): {src.name}")
            continue
        nr = f"NR-{int(m.group(1)):02d}"
        dst = DESTINO / f"{nr}.pptx"
        info = tokenizar(src, dst)
        texto = texto_completo(dst)
        resto = re.findall(r"\b\d{1,2}\s+de\s+[a-zà-ú]+\s+de\s+\d{4}\b", texto, re.IGNORECASE)
        falta = [t for t in ("{{NOME}}", "{{CPF}}") if t not in texto]
        print(f"{src.name} -> {dst.relative_to(RAIZ)}")
        print(f"    substituicoes: {info['substituicoes']} | "
              f"NOME={'ok' if '{{NOME}}' in texto else 'FALTA'} "
              f"CPF={'ok' if '{{CPF}}' in texto else 'FALTA'} "
              f"DIA={'ok' if '{{DIA}}' in texto else 'FALTA'}")
        if resto:
            ok = False
            print(f"    [AVISO] datas nao substituidas: {resto}")
        if falta:
            ok = False
            print(f"    [AVISO] tokens ausentes: {falta}")

    print("\nConcluido." if ok else "\nConcluido COM AVISOS — revise acima.")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
