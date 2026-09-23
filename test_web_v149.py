# -*- coding: utf-8 -*-
"""Testes do Portal Web — v1.48.0 (roadmap 2.35).

Padrão standalone: `python test_web_v149.py`; app FastAPI em DB tmp
(patches ANTES de create_app) + TestClient.

Cobre:
- 2.35.1 lista de abastecimentos: ordenação Data (padrão) / Serial
- 2.35.2 itens extras no form/ficha/PDF + desglose no custo (ficha e repo)
- 2.35.2 gráfico SVG de 12 meses na ficha
- 2.35.2 bloqueio/desbloqueio (motivo) + bloqueadas saem dos custos
- 2.35.2 exclusão só da solicitação mais recente
- 2.35.1 coluna Nota fiscal (número) / badge "Sem NF"
- 2.35.1 exportações: abastecimentos atualizado, custo individual e geral
"""

import io
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

FALHAS = []


def check(nome, ok):
    print(f"[{'OK' if ok else 'FALHOU'}] {nome}")
    if not ok:
        FALHAS.append(nome)


def _trocar_senha(client, atual, nova="SenhaF0rte"):
    return client.post("/troca-senha", data={"atual": atual, "nova": nova,
                                             "confirma": nova})


def make_env():
    tmp = Path(tempfile.mkdtemp(prefix="webv149_"))
    tmp.mkdir(parents=True, exist_ok=True)
    import src.utils.paths as paths_mod
    import src.core.employee_repo as er_mod
    import src.core.history_repo as hr_mod
    import src.core.frota_repo as fr_mod
    import src.core.config as config_mod
    import src.web.app as app_mod

    paths_mod.get_data_dir = lambda: tmp
    er_mod.get_db_path = lambda: tmp / "certificados.db"
    hr_mod.get_db_path = lambda: tmp / "certificados.db"
    fr_mod.get_db_path = lambda: tmp / "certificados.db"
    config_mod.load_company_config = lambda: SimpleNamespace(
        empresa_nome="Empresa Teste LTDA", empresa_cnpj="11.222.333/0001-81",
        local_treinamento="Planta Teste", instrutor_nome="Instrutor Teste",
        instrutor_registro_mte="MTE 44633/RJ")
    app_mod.get_data_dir = lambda: tmp

    from src.web.app import create_app
    from fastapi.testclient import TestClient
    app = create_app(db_path=tmp / "certificados.db",
                     secret_file=tmp / "secret.key")
    client = TestClient(app, follow_redirects=False)
    return client, tmp


def _login_admin(client, tmp):
    txt = (tmp / "web_admin_provisorio.txt").read_text(encoding="utf-8")
    prov = [l.split(":")[1].strip() for l in txt.splitlines()
            if l.startswith("SENHA PROVISORIA")][0]
    client.post("/login", data={"username": "admin", "password": prov})
    _trocar_senha(client, prov)


def main():
    client, tmp = make_env()
    _login_admin(client, tmp)

    from src.core.frota_repo import FrotaRepository
    repo = FrotaRepository()

    # ---------- dados-base ----------
    r = client.post("/frota/criar", data={
        "modelo": "Strada", "marca": "Fiat", "tipo": "pickup",
        "placa": "abd1e23", "proprio": "1", "contratante": "",
        "empresa_id": "", "obs": ""}, follow_redirects=False)
    check("V49-01 veiculo criado -> 303", r.status_code == 303)
    vid = int(r.headers["location"].split("/")[-1])

    r = client.post("/frota/fornecedores/criar", data={
        "nome": "Posto Shell", "cnpj": "33.444.555/0001-66",
        "endereco": "Av Brasil, 200"}, follow_redirects=False)
    check("V49-02 fornecedor criado -> 303", r.status_code == 303)

    # A1 (id 1): com extras + NF depois; A2 (id 2): simples; A3 (id 3): mais
    # antiga — serve para diferenciar ordenação por Data x Serial.
    r = client.post("/frota/abastecimentos/criar", data={
        "veiculo_id": str(vid), "fornecedor_id": "1", "combustivel": "gasolina",
        "data": "05/09/2026", "viagem_servico": "Obra X", "km": "10000",
        "condutor": "João", "obs": "Tanque cheio",
        "litros": "40", "valor": "289,50",
        "extra_desc_0": "Pedágio", "extra_qtd_0": "2", "extra_val_0": "30,50",
        "extra_desc_1": "", "extra_qtd_1": "", "extra_val_1": ""},
        follow_redirects=False)
    check("V49-03 A1 criado com extras -> 303", r.status_code == 303)

    r = client.post("/frota/abastecimentos/criar", data={
        "veiculo_id": str(vid), "fornecedor_id": "1", "combustivel": "gasolina",
        "data": "20/09/2026", "km": "10400", "condutor": "João",
        "litros": "20", "valor": "120,00"}, follow_redirects=False)
    check("V49-04 A2 criado -> 303", r.status_code == 303)

    r = client.post("/frota/abastecimentos/criar", data={
        "veiculo_id": str(vid), "combustivel": "outros",
        "data": "01/08/2026", "condutor": "Maria", "km": "9900"},
        follow_redirects=False)
    check("V49-05 A3 (outros, antiga) criado -> 303", r.status_code == 303)

    # ---------- 2.35.1 ordenação ----------
    r = client.get("/frota/abastecimentos")
    ok = (r.status_code == 200
          and r.text.index("AB-2026-00002") < r.text.index("AB-2026-00001")
          < r.text.index("AB-2026-00003"))
    check("V49-06 ordem padrão = Data (recentes primeiro)", ok)
    r = client.get("/frota/abastecimentos?ordem=serial")
    ok = (r.status_code == 200
          and r.text.index("AB-2026-00003") < r.text.index("AB-2026-00002")
          < r.text.index("AB-2026-00001"))
    check("V49-07 ordem=serial inverte para serial DESC", ok)

    # ---------- 2.35.1 coluna NF ----------
    r = client.get("/frota/abastecimentos")
    check("V49-08 badge Sem NF presente", 'badge b-cinza">Sem NF' in r.text)
    r = client.post("/frota/abastecimentos/1/nfs", data={
        "numero": "12345", "data_nf": "10/09/2026", "valor": "320,00"},
        files={"arquivo": ("nf.pdf", b"%PDF-1.4 teste", "application/pdf")},
        follow_redirects=False)
    check("V49-09 NF anexada -> 303", r.status_code == 303)
    r = client.get("/frota/abastecimentos")
    check("V49-10 número da NF na lista", ">12345</b>" in r.text)
    check("V49-11 Sem NF restante = 2 (A2 e A3)",
          r.text.count('badge b-cinza">Sem NF') == 2)

    # ---------- 2.35.2 extras: repo, ficha, PDF ----------
    resumo = repo.resumo_custo_veiculo(vid)
    check("V49-12 resumo.valor = combustivel+extras (440.00)",
          abs(resumo["valor"] - 440.00) < 0.01)
    check("V49-13 desglose valor_extras = 30.50",
          abs(resumo["valor_extras"] - 30.50) < 0.01)
    check("V49-14 desglose valor_combustivel = 409.50",
          abs(resumo["valor_combustivel"] - 409.50) < 0.01)

    r = client.get(f"/frota/{vid}")
    check("V49-15 ficha com desglose combustivel", "R$ 409.50" in r.text)
    check("V49-16 ficha com desglose extras", "R$ 30.50" in r.text)
    check("V49-17 gráfico SVG presente", "<svg" in r.text and "polyline" in r.text)

    r = client.get("/frota/abastecimentos/1/pdf")
    ok = r.status_code == 200 and r.content[:5] == b"%PDF-"
    texto = ""
    if ok:
        import fitz
        with fitz.open(stream=r.content, filetype="pdf") as doc:
            texto = "".join(p.get_text() for p in doc)
    check("V49-18 PDF A1 gerado", ok)
    check("V49-19 PDF com seção ITENS EXTRAS", "ITENS EXTRAS" in texto
          and "Pedágio" in texto)
    check("V49-20 PDF com TOTAL GERAL", "TOTAL GERAL" in texto)

    r = client.get("/frota/abastecimentos/1/editar")
    check("V49-21 edição pré-preenche extras", 'value="Pedágio"' in r.text
          and 'name="extra_val_0"' in r.text)
    r = client.get("/frota/abastecimentos/novo")
    check("V49-22 form novo com addExtraLinha", "addExtraLinha" in r.text
          and 'id="extras-body"' in r.text)

    # ---------- 2.35.2 bloqueio ----------
    r = client.post("/frota/abastecimentos/1/bloquear",
                    data={"motivo": "Erro de lançamento", "motivo_txt": ""},
                    follow_redirects=False)
    check("V49-23 bloquear -> 303", r.status_code == 303)
    a1 = repo.get_abastecimento(1)
    check("V49-24 status gravado", a1.get("status") == "bloqueada"
          and a1.get("motivo_status") == "Erro de lançamento")
    r = client.get("/frota/abastecimentos")
    check("V49-25 badge BLOQUEADA na lista", ">BLOQUEADA</span>" in r.text)

    resumo = repo.resumo_custo_veiculo(vid)
    check("V49-26 bloqueada sai dos totais (120.00)",
          abs(resumo["valor"] - 120.00) < 0.01)
    check("V49-27 bloqueada sai do desglose (extras 0.00)",
          resumo["valor_extras"] == 0.0)
    check("V49-28 custo_mes ignora bloqueada (120.00)",
          abs(repo.custo_mes() - 120.00) < 0.01)

    r = client.get("/frota/abastecimentos?situacao=bloqueadas")
    check("V49-29 filtro bloqueadas só A1", "AB-2026-00001" in r.text
          and "AB-2026-00002" not in r.text)
    r = client.get("/frota/abastecimentos?situacao=ativas")
    check("V49-30 filtro ativas sem A1", "AB-2026-00001" not in r.text
          and "AB-2026-00002" in r.text)

    r = client.post("/frota/abastecimentos/1/desbloquear",
                    follow_redirects=False)
    check("V49-31 desbloquear -> 303", r.status_code == 303)
    check("V49-32 desbloqueada volta ao custo (440.00)",
          abs(repo.resumo_custo_veiculo(vid)["valor"] - 440.00) < 0.01)

    # ---------- 2.35.2 exclusão (só a mais recente) ----------
    r = client.post("/frota/abastecimentos/1/excluir", follow_redirects=False)
    r2 = client.get("/frota/abastecimentos")
    check("V49-33 excluir A1 (não recente) -> 303 com orientação",
          r.status_code == 303 and "mais recente" in r2.text)
    check("V49-34 A1 continua no banco", repo.get_abastecimento(1) is not None)
    r = client.post("/frota/abastecimentos/3/excluir", follow_redirects=False)
    check("V49-35 excluir A3 (mais recente) -> 303", r.status_code == 303)
    check("V49-36 A3 removida do banco", repo.get_abastecimento(3) is None)
    r = client.get("/frota/abastecimentos")
    check("V49-37 lista sem AB-2026-00003",
          '<span class="numero">AB-2026-00003</span>' not in r.text)

    # ---------- 2.35.1 exportações ----------
    import openpyxl
    r = client.get("/frota/abastecimentos/exportar")
    ok = (r.status_code == 200
          and r.headers.get("content-type", "").startswith(
              "application/vnd.openxmlformats"))
    check("V49-38 export abastecimentos PK", ok)
    if ok:
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active
        heads = [c.value for c in ws[1]]
        check("V49-39 export com Extras/NF/Situação",
              any("extras" in str(h).lower() for h in heads)
              and any("nf" == str(h).lower() for h in heads)
              and any("situa" in str(h).lower() for h in heads))
        linhas = list(ws.iter_rows(min_row=2, values_only=True))
        check("V49-40 export linha com NF e extras",
              any("12345" in str(l) and "30.5" in str(l) for l in linhas))

    r = client.get("/frota/custos/exportar")
    ok = r.status_code == 200 and r.content[:2] == b"PK"
    check("V49-41 export custos geral PK", ok)
    if ok:
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        check("V49-42 abas Resumo por veiculo + Abastecimentos",
              "Resumo por veiculo" in wb.sheetnames
              and "Abastecimentos" in wb.sheetnames)

    r = client.get(f"/frota/{vid}/custo/exportar")
    ok = r.status_code == 200 and r.content[:2] == b"PK"
    check("V49-43 export custo individual PK", ok)
    if ok:
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        check("V49-44 aba Resumo no individual", "Resumo" in wb.sheetnames)

    print()
    if FALHAS:
        print(f"FALHAS ({len(FALHAS)}):")
        for f in FALHAS:
            print(" -", f)
        sys.exit(1)
    print("TODOS OS 44 CHECKS DE v1.48.0 PASSARAM")


if __name__ == "__main__":
    main()
